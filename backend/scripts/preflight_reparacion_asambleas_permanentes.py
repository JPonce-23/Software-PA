#!/usr/bin/env python3
"""Preflight read-only de reparación 2B-R: asambleas permanentes/continuaciones.

Este preflight atiende exclusivamente los 6 grupos que la capa 2B original
dejó en REVIEW por contener texto de "ASAMBLEA PERMANENTE"/"CONTINUACIÓN".

No altera la regla histórica general de 2B. En vez de reinterpretar cualquier
texto libre, congela los seis historiales auditados y separa explícitamente:

- Asamblea base + convocatorias -> Asamblea / AsambleaConvocatoria
- Continuación posterior -> candidato a SeguimientoEvento.continuacion_asamblea

Plan fuente congelado:

1) COYOTEPEC / ORIGEN
   Asamblea A:
     1a 2025-06-25 (programada)
     2a 2025-07-05 (programada)
   Asamblea B:
     1a 2025-09-30 (programada)
     2a 2025-10-11 (celebrada 2025-10-11)
   Continuación: 2025-10-18

2) TULA DE ALLENDE / ORIGEN
   Asamblea:
     1a 2025-08-03 (celebrada 2025-08-03)
   Continuación: 2025-08-17

3) JUCHITLAN / ORIGEN
   Asamblea:
     1a 2025-07-08 (programada)
     2a 2025-07-20 (celebrada 2025-07-20)
   Continuación: 2025-09-09

4) SAN JOSE EL MARQUEZ / ORIGEN
   Asamblea:
     1a 2025-07-11 (programada)
     2a 2025-07-20 (celebrada 2025-07-20)
   Continuación: 2025-09-09

5) SANTIAGO OXTHOC / ORIGEN
   Asamblea A:
     1a 2025-06-29 (programada)
     2a 2025-07-08 (celebrada 2025-07-08)
   Continuación: 2025-09-07
   Asamblea B:
     1a 2025-11-18 (celebrada 2025-11-18)

6) PASO DE MATA / ORIGEN
   Asamblea:
     1a 2025-08-24 (celebrada 2025-08-24)
   Continuación: 2025-08-31

Totales esperados:
- 8 Asambleas base
- 13 convocatorias
- 7 Asambleas celebradas
- 1 Asamblea sólo programada
- 6 continuaciones, que NO se escriben en esta capa

Justificación conservadora de las realizaciones que no aparecen solas en AM:
- una fecha base aparece en la misma celda temporal inmediatamente antes de
  la mención explícita "CONTINUACIÓN DE ASAMBLEA PERMANENTE";
- la continuación se separa como evento posterior, nunca como nueva Asamblea;
- Santiago Oxthoc además tiene CJ que declara la sesión permanente y soporte
  de 2a convocatoria del 08/07;
- Paso de Mata declara explícitamente la sesión permanente tras 24/08.

Este script NO escribe en PostgreSQL.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.database import DB_NAME, SessionLocal  # noqa: E402


EXPECTED_DB = "db_carga_excel"
EXPECTED_SCHEMA = "018"
EXPECTED_SHA256 = "bac01c25e9ff963f524fa8a48d69e74d92a0d36578731b1eca91f0e5a3573c3f"

PROJECT_KEY = "MEX-QRO"
SHEET = "INFORME M-Q"

COP_CONTEXT = {
    "ORIGEN": "cop_original",
    "ADICIONAL": "modificatorio",
    "2A_ADICIONAL": "modificatorio",
    "COMPLEMENTARIAS": "obras_complementarias",
}

EXPECTED_ASSEMBLIES = 8
EXPECTED_CONVOCATIONS = 13
EXPECTED_CELEBRATED = 7
EXPECTED_PROGRAMMED_ONLY = 1
EXPECTED_CONTINUATIONS = 6


class PreflightAbort(RuntimeError):
    pass


def d(value: str) -> date:
    return date.fromisoformat(value)


PLAN = [
    {
        "source_key": "MEXICO|COYOTEPEC|COYOTEPEC",
        "cop": "ORIGEN",
        "source_rows": [45, 47, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59],
        "source_checks": {
            "AI_dates": {d("2025-06-25"), d("2025-09-30")},
            "AJ_dates": {d("2025-07-05"), d("2025-10-11"), d("2025-10-18")},
            "AM_dates": {d("2025-10-11"), d("2025-10-18")},
            "must_text": ["ASAMBLEA PERMANENTE", "CONTINUACION"],
        },
        "assemblies": [
            {
                "label": "COYOTEPEC_A",
                "rule": "ciclo_previo_programado",
                "convocations": [
                    (1, d("2025-06-25"), None, None),
                    (2, d("2025-07-05"), None, None),
                ],
            },
            {
                "label": "COYOTEPEC_B",
                "rule": "base_permanente_segunda_convocatoria",
                "convocations": [
                    (1, d("2025-09-30"), None, None),
                    (2, d("2025-10-11"), d("2025-10-11"), "celebrada"),
                ],
            },
        ],
        "continuation": {
            "assembly_label": "COYOTEPEC_B",
            "date": d("2025-10-18"),
            "source_column": "CJ",
            "source_row": 45,
        },
    },
    {
        "source_key": "HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE",
        "cop": "ORIGEN",
        "source_rows": [92, 93],
        "source_checks": {
            "AI_dates": {d("2025-08-03")},
            "AJ_dates": set(),
            "AM_dates": {d("2025-08-03"), d("2025-08-17")},
            "must_text": ["CONTINUACION DE ASAMBLEA"],
        },
        "assemblies": [
            {
                "label": "TULA_A",
                "rule": "base_permanente_primera_convocatoria",
                "convocations": [
                    (1, d("2025-08-03"), d("2025-08-03"), "celebrada"),
                ],
            },
        ],
        "continuation": {
            "assembly_label": "TULA_A",
            "date": d("2025-08-17"),
            "source_column": "CJ",
            "source_row": 92,
        },
    },
    {
        "source_key": "HIDALGO|CHAPANTONGO|JUCHITLAN",
        "cop": "ORIGEN",
        "source_rows": [103],
        "source_checks": {
            "AI_dates": {d("2025-07-08")},
            "AJ_dates": {d("2025-07-20"), d("2025-09-09")},
            "AM_dates": {d("2025-07-20"), d("2025-09-09")},
            "must_text": ["CONTINUACION DE ASAMBLEA PERMANENTE"],
        },
        "assemblies": [
            {
                "label": "JUCHITLAN_A",
                "rule": "base_permanente_segunda_convocatoria",
                "convocations": [
                    (1, d("2025-07-08"), None, None),
                    (2, d("2025-07-20"), d("2025-07-20"), "celebrada"),
                ],
            },
        ],
        "continuation": {
            "assembly_label": "JUCHITLAN_A",
            "date": d("2025-09-09"),
            "source_column": "AJ",
            "source_row": 103,
        },
    },
    {
        "source_key": "HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ",
        "cop": "ORIGEN",
        "source_rows": [105, 106],
        "source_checks": {
            "AI_dates": {d("2025-07-11")},
            "AJ_dates": {d("2025-07-20")},
            "AM_dates": {d("2025-07-20"), d("2025-09-09")},
            "must_text": ["CONTINUACION DE ASAMBLEA PERMANENTE"],
        },
        "assemblies": [
            {
                "label": "SAN_JOSE_A",
                "rule": "base_permanente_segunda_convocatoria",
                "convocations": [
                    (1, d("2025-07-11"), None, None),
                    (2, d("2025-07-20"), d("2025-07-20"), "celebrada"),
                ],
            },
        ],
        "continuation": {
            "assembly_label": "SAN_JOSE_A",
            "date": d("2025-09-09"),
            "source_column": "AJ",
            "source_row": 105,
        },
    },
    {
        "source_key": "MEXICO|JILOTEPEC|SANTIAGO OXTHOC",
        "cop": "ORIGEN",
        "source_rows": [117],
        "source_checks": {
            "AI_dates": {d("2025-06-29"), d("2025-11-18")},
            "AJ_dates": {d("2025-07-08"), d("2025-09-07")},
            "AM_dates": {d("2025-11-18")},
            "must_text": ["CONTINUACION DE ASAMBLEA PERMANENTE", "SE DECLARA ASAMBLEA PERMANENTE"],
        },
        "assemblies": [
            {
                "label": "SANTIAGO_A",
                "rule": "base_permanente_segunda_convocatoria_textual",
                "convocations": [
                    (1, d("2025-06-29"), None, None),
                    (2, d("2025-07-08"), d("2025-07-08"), "celebrada"),
                ],
            },
            {
                "label": "SANTIAGO_B",
                "rule": "asamblea_posterior_primera_convocatoria",
                "convocations": [
                    (1, d("2025-11-18"), d("2025-11-18"), "celebrada"),
                ],
            },
        ],
        "continuation": {
            "assembly_label": "SANTIAGO_A",
            "date": d("2025-09-07"),
            "source_column": "AJ",
            "source_row": 117,
        },
    },
    {
        "source_key": "QUERETARO|SAN JUAN DEL RIO|PASO DE MATA",
        "cop": "ORIGEN",
        "source_rows": [136, 138, 139],
        "source_checks": {
            "AI_dates": {d("2025-08-24")},
            "AJ_dates": set(),
            "AM_dates": {d("2025-08-31")},
            "must_text": ["SE DECLARO ASAMBLEA PERMANENTE", "CONTINUACION"],
        },
        "assemblies": [
            {
                "label": "PASO_MATA_A",
                "rule": "base_permanente_primera_convocatoria_textual",
                "convocations": [
                    (1, d("2025-08-24"), d("2025-08-24"), "celebrada"),
                ],
            },
        ],
        "continuation": {
            "assembly_label": "PASO_MATA_A",
            "date": d("2025-08-31"),
            "source_column": "AI",
            "source_row": 136,
        },
    },
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    value = str(value).replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    return value.strip()


def norm(value: Any) -> str:
    raw = clean_text(value).upper()
    normalized = unicodedata.normalize("NFKD", raw)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9]+", " ", ascii_text).strip()


def make_nucleus_key(entity: Any, municipality: Any, nucleus: Any) -> str | None:
    parts = (norm(entity), norm(municipality), norm(nucleus))
    return "|".join(parts) if all(parts) else None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def schema_version(db) -> str | None:
    return db.execute(text("SELECT max(version) FROM schema_migrations")).scalar()


DATE_PATTERN = re.compile(r"(?<!\d)(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})(?!\d)")

MONTHS = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SEPTIEMBRE": 9,
    "SETIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}

MONTH_DATE_PATTERN = re.compile(
    r"(?<!\d)(\d{1,2})\s+(?:DE\s+)?"
    r"(ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|"
    r"SEPTIEMBRE|SETIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE)"
    r"(?:\s+DE)?\s+(\d{4})(?!\d)",
    re.IGNORECASE,
)


def parse_dates(value: Any) -> list[date]:
    """Extrae sólo fechas con año explícito.

    Acepta tanto 20/07/2025 como '9 DE SEPTIEMBRE DE 2025'.
    Deliberadamente NO completa años ausentes, por ejemplo
    '31 DE AGOSTO', porque eso pertenece a la corroboración posterior.
    """
    if value is None:
        return []
    if isinstance(value, datetime):
        return [value.date()]
    if isinstance(value, date):
        return [value]

    raw = clean_text(value)
    found: list[date] = []

    for day_s, month_s, year_s in DATE_PATTERN.findall(raw):
        year = int(year_s)
        if year < 100:
            year += 2000
        try:
            candidate = date(year, int(month_s), int(day_s))
        except ValueError:
            continue
        if candidate not in found:
            found.append(candidate)

    normalized = unicodedata.normalize("NFKD", raw.upper())
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")

    for day_s, month_name, year_s in MONTH_DATE_PATTERN.findall(ascii_text):
        try:
            candidate = date(
                int(year_s),
                MONTHS[month_name.upper()],
                int(day_s),
            )
        except (ValueError, KeyError):
            continue
        if candidate not in found:
            found.append(candidate)

    return found


def catalog_by_code(db, catalog_type: str) -> dict[str, Any]:
    rows = (
        db.query(models.CatalogoOperativo)
        .filter(
            models.CatalogoOperativo.tipo_catalogo == catalog_type,
            models.CatalogoOperativo.activo.is_(True),
        )
        .all()
    )
    return {row.codigo: row for row in rows}


def build_project_nucleus_index(db, project):
    rows = (
        db.query(
            models.ProyectoNucleo,
            models.NucleoAgrario,
            models.Municipio,
            models.EntidadFederativa,
        )
        .join(
            models.NucleoAgrario,
            models.NucleoAgrario.id_nucleo == models.ProyectoNucleo.id_nucleo,
        )
        .join(
            models.Municipio,
            models.Municipio.id_municipio == models.NucleoAgrario.id_municipio,
        )
        .join(
            models.EntidadFederativa,
            models.EntidadFederativa.id_entidad == models.Municipio.id_entidad,
        )
        .filter(
            models.ProyectoNucleo.id_proyecto == project.id_proyecto,
            models.ProyectoNucleo.activo.is_(True),
            models.NucleoAgrario.activo.is_(True),
        )
        .all()
    )
    index = {}
    for pn, nucleus, municipality, entity in rows:
        key = make_nucleus_key(entity.nombre, municipality.nombre, nucleus.nombre_nucleo)
        if key in index:
            raise PreflightAbort(f"ProyectoNucleo duplicado para {key}")
        index[key] = pn
    return index


def candidate_signature(candidate: dict[str, Any]) -> tuple:
    return tuple(candidate["convocations"])


def existing_assemblies(db, pn_id, type_id, context_id, cop_id, result_catalog):
    result_by_id = {
        item.id_catalogo_opcion: item.codigo
        for item in result_catalog.values()
    }

    assemblies = (
        db.query(models.Asamblea)
        .filter(
            models.Asamblea.id_proyecto_nucleo == pn_id,
            models.Asamblea.id_tipo_asamblea == type_id,
            models.Asamblea.id_contexto_asamblea == context_id,
            models.Asamblea.id_tipo_cop_operativo == cop_id,
            models.Asamblea.activo.is_(True),
        )
        .order_by(models.Asamblea.id_asamblea)
        .all()
    )

    output = []
    for assembly in assemblies:
        convs = (
            db.query(models.AsambleaConvocatoria)
            .filter(
                models.AsambleaConvocatoria.id_asamblea == assembly.id_asamblea,
                models.AsambleaConvocatoria.activo.is_(True),
            )
            .order_by(models.AsambleaConvocatoria.ordinal)
            .all()
        )
        signature = tuple(
            (
                conv.ordinal,
                conv.fecha_programada,
                conv.fecha_realizacion,
                result_by_id.get(conv.id_resultado),
            )
            for conv in convs
        )
        output.append((assembly, signature))
    return output


def validate_source(sheet, group):
    row0 = group["source_rows"][0]
    key = make_nucleus_key(
        sheet[f"B{row0}"].value,
        sheet[f"C{row0}"].value,
        sheet[f"F{row0}"].value,
    )
    if key != group["source_key"]:
        raise PreflightAbort(
            f"Clave fuente cambió en fila {row0}: esperado={group['source_key']}, actual={key}"
        )

    cop = clean_text(sheet[f"M{row0}"].value).upper().replace(" ", "_")
    if cop != group["cop"]:
        raise PreflightAbort(
            f"COP cambió para {group['source_key']}: esperado={group['cop']}, actual={cop}"
        )

    observed = {
        "AI_dates": set(parse_dates(sheet[f"AI{row0}"].value)),
        "AJ_dates": set(parse_dates(sheet[f"AJ{row0}"].value)),
        "AM_dates": set(parse_dates(sheet[f"AM{row0}"].value)),
    }

    # Algunas variantes repetidas aparecen en filas adicionales; unimos AM de
    # todas las filas del grupo para capturar valores ISO duplicados.
    observed["AM_dates"] = set()
    combined_text_parts = []
    for row in group["source_rows"]:
        observed["AM_dates"].update(parse_dates(sheet[f"AM{row}"].value))
        for col in ("AI", "AJ", "AM", "CJ"):
            value = clean_text(sheet[f"{col}{row}"].value)
            if value:
                combined_text_parts.append(value)

    combined_text = norm("\n".join(combined_text_parts))

    for field in ("AI_dates", "AJ_dates"):
        expected = group["source_checks"][field]
        if observed[field] != expected:
            raise PreflightAbort(
                f"{group['source_key']} {field} cambió: "
                f"esperado={sorted(expected)}, actual={sorted(observed[field])}"
            )

    expected_am = group["source_checks"]["AM_dates"]
    if observed["AM_dates"] != expected_am:
        raise PreflightAbort(
            f"{group['source_key']} AM_dates cambió: "
            f"esperado={sorted(expected_am)}, actual={sorted(observed['AM_dates'])}"
        )

    for marker in group["source_checks"]["must_text"]:
        if norm(marker) not in combined_text:
            raise PreflightAbort(
                f"{group['source_key']} perdió marcador textual esperado: {marker!r}"
            )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    args = parser.parse_args()

    if DB_NAME != EXPECTED_DB:
        raise PreflightAbort(
            f"Protección: DB_NAME debe ser {EXPECTED_DB!r}; recibido {DB_NAME!r}."
        )

    if not args.excel.exists():
        raise PreflightAbort(f"No existe el Excel: {args.excel}")

    digest = sha256_file(args.excel)
    if digest != EXPECTED_SHA256:
        raise PreflightAbort(
            "El SHA-256 no coincide con el Excel aprobado. "
            f"Esperado {EXPECTED_SHA256}; recibido {digest}."
        )

    workbook = load_workbook(args.excel, data_only=True, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise PreflightAbort(f"Falta la hoja {SHEET!r}.")
    sheet = workbook[SHEET]

    for group in PLAN:
        validate_source(sheet, group)

    expected_assemblies = sum(len(group["assemblies"]) for group in PLAN)
    expected_convs = sum(
        len(assembly["convocations"])
        for group in PLAN
        for assembly in group["assemblies"]
    )
    expected_celebrated = sum(
        1
        for group in PLAN
        for assembly in group["assemblies"]
        if any(conv[3] == "celebrada" for conv in assembly["convocations"])
    )

    if expected_assemblies != EXPECTED_ASSEMBLIES:
        raise PreflightAbort("Plan interno de Asambleas cambió.")
    if expected_convs != EXPECTED_CONVOCATIONS:
        raise PreflightAbort("Plan interno de convocatorias cambió.")
    if expected_celebrated != EXPECTED_CELEBRATED:
        raise PreflightAbort("Plan interno de celebradas cambió.")
    if len(PLAN) != EXPECTED_CONTINUATIONS:
        raise PreflightAbort("Plan interno de continuaciones cambió.")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.output_dir / "preflight_reparacion_asambleas_permanentes.csv"
    json_path = args.output_dir / "preflight_reparacion_asambleas_permanentes.json"

    db = SessionLocal()
    try:
        current_db = db.execute(text("SELECT current_database()")).scalar()
        current_schema = schema_version(db)
        if current_db != EXPECTED_DB or current_schema != EXPECTED_SCHEMA:
            raise PreflightAbort(
                f"Destino inesperado: BD={current_db!r}, schema={current_schema!r}."
            )

        project = (
            db.query(models.Proyecto)
            .filter(
                models.Proyecto.clave_proyecto == PROJECT_KEY,
                models.Proyecto.activo.is_(True),
            )
            .one_or_none()
        )
        if project is None:
            raise PreflightAbort(f"No existe proyecto activo {PROJECT_KEY!r}.")

        pn_index = build_project_nucleus_index(db, project)
        cop_catalog = catalog_by_code(db, "tipo_cop_operativo")
        type_catalog = catalog_by_code(db, "tipo_asamblea")
        context_catalog = catalog_by_code(db, "contexto_asamblea")
        result_catalog = catalog_by_code(db, "resultado_convocatoria")

        for code in ("anuencia",):
            if code not in type_catalog:
                raise PreflightAbort(f"Falta tipo_asamblea/{code}.")
        if "celebrada" not in result_catalog:
            raise PreflightAbort("Falta resultado_convocatoria/celebrada.")

        actions = Counter()
        conv_actions = Counter()
        details = []
        continuation_rows = []

        for group in PLAN:
            pn = pn_index.get(group["source_key"])
            if pn is None:
                raise PreflightAbort(
                    f"ProyectoNucleo faltante para grupo reparable: {group['source_key']}"
                )

            cop = cop_catalog.get(group["cop"])
            if cop is None:
                raise PreflightAbort(f"Falta tipo_cop_operativo/{group['cop']}.")

            context_code = COP_CONTEXT[group["cop"]]
            context = context_catalog.get(context_code)
            if context is None:
                raise PreflightAbort(f"Falta contexto_asamblea/{context_code}.")

            existing = existing_assemblies(
                db,
                pn.id_proyecto_nucleo,
                type_catalog["anuencia"].id_catalogo_opcion,
                context.id_catalogo_opcion,
                cop.id_catalogo_opcion,
                result_catalog,
            )

            matched_labels = {}

            for candidate in group["assemblies"]:
                signature = candidate_signature(candidate)
                exact = [
                    assembly
                    for assembly, actual_signature in existing
                    if actual_signature == signature
                ]

                if len(exact) > 1:
                    action = "REVIEW"
                    detail = (
                        f"Hay {len(exact)} Asambleas exactas activas con la misma firma."
                    )
                    assembly_id = None
                elif len(exact) == 1:
                    action = "REUSE"
                    detail = f"Reutilizar id_asamblea={exact[0].id_asamblea}."
                    assembly_id = exact[0].id_asamblea
                    matched_labels[candidate["label"]] = assembly_id
                else:
                    # Si existe cualquier otra asamblea del mismo PN+COP, no la
                    # pisamos; el importador posterior deberá volver a validar.
                    action = "CREATE"
                    detail = "Asamblea permanente/posterior determinista candidata."
                    assembly_id = None

                actions[action] += 1
                conv_actions["REUSE" if action == "REUSE" else "CREATE"] += (
                    len(candidate["convocations"]) if action != "REVIEW" else 0
                )

                details.append({
                    "action": action,
                    "source_key": group["source_key"],
                    "source_rows": ";".join(map(str, group["source_rows"])),
                    "cop": group["cop"],
                    "label": candidate["label"],
                    "rule": candidate["rule"],
                    "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                    "id_asamblea": assembly_id or "",
                    "convocatorias": " | ".join(
                        (
                            f"{ordinal}a:"
                            f"prog={programmed.isoformat()},"
                            f"real={realized.isoformat() if realized else '-'},"
                            f"resultado={result or '-'}"
                        )
                        for ordinal, programmed, realized, result
                        in candidate["convocations"]
                    ),
                    "detail": detail,
                })

            continuation = group["continuation"]
            continuation_rows.append({
                "source_key": group["source_key"],
                "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                "cop": group["cop"],
                "assembly_label": continuation["assembly_label"],
                "continuation_date": continuation["date"].isoformat(),
                "source_column": continuation["source_column"],
                "source_row": continuation["source_row"],
                "status": (
                    "TARGET_REUSE_READY"
                    if continuation["assembly_label"] in matched_labels
                    else "TARGET_AFTER_REPAIR"
                ),
            })

        if actions["REVIEW"] != 0:
            # El reporte se genera, pero la salida deja claro que no es plan
            # limpio para escritura.
            pass

        with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(
                stream,
                fieldnames=[
                    "action",
                    "source_key",
                    "source_rows",
                    "cop",
                    "label",
                    "rule",
                    "id_proyecto_nucleo",
                    "id_asamblea",
                    "convocatorias",
                    "detail",
                ],
            )
            writer.writeheader()
            writer.writerows(details)

        summary = {
            "database": current_db,
            "schema": current_schema,
            "project": PROJECT_KEY,
            "excel": args.excel.name,
            "sha256": digest,
            "plan": {
                "groups": len(PLAN),
                "assemblies": EXPECTED_ASSEMBLIES,
                "convocations": EXPECTED_CONVOCATIONS,
                "celebrated": EXPECTED_CELEBRATED,
                "programmed_only": EXPECTED_PROGRAMMED_ONLY,
                "continuations": EXPECTED_CONTINUATIONS,
            },
            "actions": {
                code: actions[code]
                for code in ("CREATE", "REUSE", "REVIEW")
            },
            "continuations": continuation_rows,
        }

        json_path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        print("=== PREFLIGHT REPARACIÓN 2B-R — ASAMBLEAS PERMANENTES ===")
        print(f"BD: {current_db} / schema {current_schema}")
        print(f"Proyecto: {PROJECT_KEY}")
        print(f"SHA-256: {digest}")
        print(f"Grupos fuente:              {len(PLAN)}")
        print(f"Asambleas base:             {EXPECTED_ASSEMBLIES}")
        print(f"Convocatorias:              {EXPECTED_CONVOCATIONS}")
        print(f"Asambleas celebradas:       {EXPECTED_CELEBRATED}")
        print(f"Asambleas sólo programadas: {EXPECTED_PROGRAMMED_ONLY}")
        print(f"Continuaciones separadas:   {EXPECTED_CONTINUATIONS}")
        print("Acciones Asambleas:")
        for code in ("CREATE", "REUSE", "REVIEW"):
            print(f"  {code:<8} {actions[code]}")

        print("Continuaciones (no se escriben aquí):")
        for item in continuation_rows:
            print(
                f"  {item['source_key']} | "
                f"{item['continuation_date']} | "
                f"{item['assembly_label']} | "
                f"{item['status']}"
            )

        print(
            "Nota: las continuaciones se migrarán después como "
            "SeguimientoEvento.continuacion_asamblea."
        )
        print(
            "Nota: este preflight sólo repara la estructura que 2B dejó en REVIEW."
        )
        print(f"Reporte JSON: {json_path}")
        print(f"Plan CSV:     {csv_path}")
        print("No se realizó ninguna escritura en PostgreSQL.")
        return 0

    finally:
        db.rollback()
        db.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except PreflightAbort as exc:
        print("=== PREFLIGHT ABORTADO ===", file=sys.stderr)
        print(str(exc), file=sys.stderr)
        raise SystemExit(2)
