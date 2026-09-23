#!/usr/bin/env python3
"""Resolución read-only 2J-R posterior a 2B-R/2F-R/2I-B-R.

Problema corregido
==================
El resolver original de 2J:
- analiza correctamente 38 EVENT_CANDIDATE en CG/CJ;
- obtiene 13 eventos con fecha determinista y 25 REVIEW_DATE;
- pero para continuacion_asamblea buscaba una Asamblea cuya convocatoria
  celebrada tuviera la MISMA fecha que la continuación.

Después de 2B-R esa regla ya no es válida:
la continuación es un SeguimientoEvento posterior que apunta a la Asamblea
base, no una nueva Asamblea celebrada en la fecha de continuación.

Además, 2B-R auditó seis continuaciones en total:
- 2 ya están presentes como EVENT_CANDIDATE en CJ (Coyotepec y Tula);
- 4 están en AI/AJ y no forman parte de las 183 cláusulas CG/CJ
  (Juchitlán, San José el Márquez, Santiago Oxthoc y Paso de Mata).

Este resolver:
1) conserva sin reinterpretar las 38 candidatas originales de CG/CJ;
2) conserva el corte temporal original: 13 deterministas + 25 REVIEW_DATE;
3) resuelve las 2 continuaciones de CJ contra la Asamblea base por la firma
   semántica congelada de 2B-R;
4) añade exactamente las 4 continuaciones de AI/AJ auditadas por 2B-R;
5) valida la fuente 2B-R contra el Excel antes de usar esas cuatro;
6) no congela IDs de Asamblea;
7) compara con SeguimientoEvento para CREATE/REUSE exacto;
8) NO escribe en PostgreSQL.

Plan final esperado antes de importar 2J:
- CG/CJ: 38 EVENT_CANDIDATE
- REVIEW_DATE: 25
- deterministas CG/CJ: 13
- continuaciones extra AI/AJ: 4
- deterministas finales: 17
    reunion: 8
    cambio_alcance/nueva_informacion: 2
    medicion_bdt: 1
    continuacion_asamblea: 6
- BLOCKED_TARGET: 0
- REVIEW_TARGET: 0
- SKIP_SOURCE: 0
- REVIEW_EXISTING: 0
- en estado previo a 2J: CREATE_CANDIDATE=17 / REUSE_CANDIDATE=0
- después de importar 2J: CREATE_CANDIDATE=0 / REUSE_CANDIDATE=17
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.database import DB_NAME, SessionLocal  # noqa: E402

import resolver_seguimiento_cg_cj as base  # noqa: E402
import importar_reparacion_asambleas_permanentes as repair2br  # noqa: E402
import preflight_reparacion_asambleas_permanentes as source2br  # noqa: E402


EXPECTED_DB = "db_carga_excel"
EXPECTED_SCHEMA = "018"
EXPECTED_SHA256 = "bac01c25e9ff963f524fa8a48d69e74d92a0d36578731b1eca91f0e5a3573c3f"
PROJECT_KEY = "MEX-QRO"
SHEET = "INFORME M-Q"

EXPECTED_ANALYSIS_ROWS = 183
EXPECTED_EVENT_CANDIDATES = 38
EXPECTED_REVIEW_DATE = 25
EXPECTED_CG_CJ_DETERMINISTIC = 13
EXPECTED_EXTRA_CONTINUATIONS = 4
EXPECTED_FINAL_DETERMINISTIC = 17

EXPECTED_FINAL_BY_EVENT = {
    ("cambio_alcance", "nueva_informacion"): 2,
    ("continuacion_asamblea", ""): 6,
    ("medicion_bdt", ""): 1,
    ("reunion", ""): 8,
}

# Las cuatro continuaciones que 2B-R detectó fuera de CG/CJ.
EXPECTED_EXTRA_KEYS = {
    ("HIDALGO|CHAPANTONGO|JUCHITLAN", "AJ", 103),
    ("HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ", "AJ", 105),
    ("MEXICO|JILOTEPEC|SANTIAGO OXTHOC", "AJ", 117),
    ("QUERETARO|SAN JUAN DEL RIO|PASO DE MATA", "AI", 136),
}


class ResolveAbort(RuntimeError):
    pass


def catalog_by_code(db, catalog_type: str) -> dict[str, Any]:
    return base.catalog_by_code(db, catalog_type)


def parse_excel_full_dates(value: Any) -> list[date]:
    """Parsea una fecha estructurada de Excel sin reinterpretarla.

    openpyxl entrega las celdas con formato fecha como datetime/date. El
    parser textual histórico de 2J convierte esos objetos a ISO YYYY-MM-DD
    pero después busca DD/MM/YYYY, por lo que puede perder una fecha válida.
    Para objetos estructurados se conserva directamente la fecha; para texto
    se reutiliza el parser histórico.
    """
    if value is None:
        return []
    if isinstance(value, datetime):
        return [value.date()]
    if isinstance(value, date):
        return [value]
    return base.parse_full_dates(value)


def corroborate_continuation_date_safe(
    sheet,
    row: int,
    clause: str,
) -> tuple[date | None, str]:
    """Corrobora día/mes de continuación contra AM sin inferir el año."""
    partials = base.parse_partial_day_month(clause)

    if len(partials) != 1:
        return None, (
            "Continuación: se esperó un único día/mes sin año; "
            f"se detectaron {len(partials)}."
        )

    structured_dates = parse_excel_full_dates(sheet[f"AM{row}"].value)
    matches = [
        value
        for value in structured_dates
        if (value.day, value.month) == partials[0]
    ]

    if len(matches) != 1:
        return None, (
            "Continuación: el día/mes no resolvió de forma única contra AM; "
            f"coincidencias={len(matches)}."
        )

    return matches[0], (
        "Fecha corroborada por día/mes contra fecha estructurada AM "
        f"de la misma fila: {matches[0].isoformat()}."
    )


def exact_existing_event(
    db,
    *,
    pn_id: int,
    target_type: str,
    target_id: int,
    event_type_id: int,
    motive_id: int | None,
    event_date,
    detail: str,
    source: str,
):
    query = db.query(models.SeguimientoEvento).filter(
        models.SeguimientoEvento.id_proyecto_nucleo == pn_id,
        models.SeguimientoEvento.entidad_tipo == target_type,
        models.SeguimientoEvento.entidad_id == target_id,
        models.SeguimientoEvento.ambito == "colectivo",
        models.SeguimientoEvento.id_tipo_evento == event_type_id,
        models.SeguimientoEvento.fecha_evento == event_date,
        models.SeguimientoEvento.detalle == detail,
        models.SeguimientoEvento.id_documento.is_(None),
        models.SeguimientoEvento.fuente == source,
        models.SeguimientoEvento.activo.is_(True),
    )
    if motive_id is None:
        query = query.filter(models.SeguimientoEvento.id_motivo.is_(None))
    else:
        query = query.filter(models.SeguimientoEvento.id_motivo == motive_id)

    return query.order_by(models.SeguimientoEvento.id_seguimiento_evento).all()


def group_for_source_key(source_key: str) -> dict[str, Any]:
    matches = [
        group
        for group in repair2br.PLAN
        if group["source_key"] == source_key
    ]
    if len(matches) != 1:
        raise ResolveAbort(
            f"2B-R no contiene exactamente un grupo para {source_key}: "
            f"{len(matches)}."
        )
    return matches[0]


def resolve_2br_base_assembly(
    db,
    *,
    source_key: str,
    pn,
    cop_catalog,
    context_catalog,
    assembly_type_catalog,
    result_catalog,
):
    """Resuelve la Asamblea base por la firma exacta congelada en 2B-R."""
    group = group_for_source_key(source_key)
    continuation = group["continuation"]

    if "anuencia" not in assembly_type_catalog:
        raise ResolveAbort("Falta tipo_asamblea/anuencia.")

    cop = cop_catalog.get(group["cop"])
    if cop is None:
        raise ResolveAbort(
            f"Falta tipo_cop_operativo/{group['cop']} para {source_key}."
        )

    context_code = repair2br.COP_CONTEXT[group["cop"]]
    context = context_catalog.get(context_code)
    if context is None:
        raise ResolveAbort(
            f"Falta contexto_asamblea/{context_code} para {source_key}."
        )

    result_by_id = {
        item.id_catalogo_opcion: item.codigo
        for item in result_catalog.values()
    }

    assemblies = repair2br.active_group_assemblies(
        db,
        pn.id_proyecto_nucleo,
        assembly_type_catalog["anuencia"].id_catalogo_opcion,
        context.id_catalogo_opcion,
        cop.id_catalogo_opcion,
        result_by_id,
    )

    candidate_matches = [
        candidate
        for candidate in group["assemblies"]
        if candidate["label"] == continuation["assembly_label"]
    ]
    if len(candidate_matches) != 1:
        raise ResolveAbort(
            "No se resolvió de forma única el assembly_label congelado: "
            f"{source_key}/{continuation['assembly_label']}."
        )

    signature = repair2br.candidate_signature(candidate_matches[0])
    exact = [
        assembly
        for assembly, actual_signature in assemblies
        if actual_signature == signature
    ]

    if len(exact) != 1:
        raise ResolveAbort(
            "La Asamblea base 2B-R no resuelve por firma exacta: "
            f"{source_key}/{continuation['assembly_label']}; "
            f"coincidencias={len(exact)}."
        )

    return exact[0], group


def continuation_fragment(value: Any) -> str:
    """Extrae del texto AI/AJ el fragmento explícito de continuación.

    No inventa contenido. Si la fecha está en la línea inmediata posterior
    (Santiago Oxthoc), la concatena al fragmento.
    """
    raw = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in raw.split("\n") if line.strip()]

    for index, line in enumerate(lines):
        if "CONTINU" not in base.norm(line):
            continue

        parts = [line]
        # Caso auditado: "EL DÍA" en una línea y fecha en la siguiente.
        if index + 1 < len(lines):
            next_line = lines[index + 1]
            if base.parse_full_dates(next_line):
                if not base.parse_full_dates(line):
                    parts.append(next_line)

        return base.clean_text(" ".join(parts))

    raise ResolveAbort("La celda 2B-R no contiene fragmento de continuación.")


def extra_source_string(row: int, column: str) -> str:
    return (
        f"Excel INFORME M-Q!{column} fila {row} "
        "continuación auditada 2B-R"
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument(
        "--analysis-csv",
        type=Path,
        default=Path("/data/reportes/analisis_seguimiento_cg_cj_v2.csv"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("/data/reportes/2j_post_reparaciones"),
    )
    args = parser.parse_args()

    if DB_NAME != EXPECTED_DB:
        raise ResolveAbort(
            f"Protección: DB_NAME debe ser {EXPECTED_DB!r}; recibido {DB_NAME!r}."
        )
    if not args.excel.exists():
        raise ResolveAbort(f"No existe el Excel: {args.excel}")
    if not args.analysis_csv.exists():
        raise ResolveAbort(
            f"No existe el CSV de análisis: {args.analysis_csv}."
        )

    digest = base.sha256_file(args.excel)
    if digest != EXPECTED_SHA256:
        raise ResolveAbort(
            "El SHA-256 no coincide con el Excel aprobado. "
            f"Esperado {EXPECTED_SHA256}; recibido {digest}."
        )

    with args.analysis_csv.open(
        "r",
        newline="",
        encoding="utf-8-sig",
    ) as stream:
        analysis_rows = list(csv.DictReader(stream))

    if len(analysis_rows) != EXPECTED_ANALYSIS_ROWS:
        raise ResolveAbort(
            f"Cambió análisis CG/CJ: esperado={EXPECTED_ANALYSIS_ROWS}, "
            f"actual={len(analysis_rows)}."
        )

    candidates = [
        row
        for row in analysis_rows
        if row["classification"].strip() == "EVENT_CANDIDATE"
    ]
    if len(candidates) != EXPECTED_EVENT_CANDIDATES:
        raise ResolveAbort(
            f"Cambió EVENT_CANDIDATE: esperado={EXPECTED_EVENT_CANDIDATES}, "
            f"actual={len(candidates)}."
        )

    # Se conservan exactamente las guardas semánticas del análisis original.
    event_counts = Counter(
        (row["event_type"].strip(), row["motive"].strip())
        for row in candidates
    )
    for key, expected in base.EXPECTED_EVENT_COUNTS.items():
        if event_counts[key] != expected:
            raise ResolveAbort(
                f"Cambió candidato CG/CJ {key}: "
                f"esperado={expected}, actual={event_counts[key]}."
            )

    workbook = load_workbook(args.excel, data_only=True, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise ResolveAbort(f"Falta la hoja {SHEET!r}.")
    sheet = workbook[SHEET]

    # Revalidar las seis fuentes auditadas de 2B-R contra el Excel antes de
    # usarlas como plan de continuidad.
    for group in source2br.PLAN:
        source2br.validate_source(sheet, group)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.output_dir / "resolucion_seguimiento_post_reparaciones.csv"
    json_path = args.output_dir / "resolucion_seguimiento_post_reparaciones.json"

    db = SessionLocal()
    try:
        current_db = db.execute(text("SELECT current_database()")).scalar()
        current_schema = base.schema_version(db)
        if current_db != EXPECTED_DB or current_schema != EXPECTED_SCHEMA:
            raise ResolveAbort(
                f"Destino inesperado: BD={current_db!r}, schema={current_schema!r}."
            )

        project = (
            db.query(models.Proyecto)
            .filter(
                models.Proyecto.clave_proyecto == PROJECT_KEY,
                models.Proyecto.activo.is_(True),
            )
            .one_or_none()
        )
        if project is None:
            raise ResolveAbort(f"No existe proyecto activo {PROJECT_KEY!r}.")

        pn_index = base.build_project_nucleus_index(db, project)
        event_catalog = catalog_by_code(db, "tipo_evento_seguimiento")
        motive_catalog = catalog_by_code(db, "motivo_seguimiento")
        cop_catalog = catalog_by_code(db, "tipo_cop_operativo")
        context_catalog = catalog_by_code(db, "contexto_asamblea")
        assembly_type_catalog = catalog_by_code(db, "tipo_asamblea")
        result_catalog = catalog_by_code(db, "resultado_convocatoria")

        required_events = {
            "reunion",
            "cambio_alcance",
            "medicion_bdt",
            "continuacion_asamblea",
        }
        missing_events = required_events - set(event_catalog)
        if missing_events:
            raise ResolveAbort(
                "Faltan tipos de seguimiento: "
                + ", ".join(sorted(missing_events))
            )
        if "nueva_informacion" not in motive_catalog:
            raise ResolveAbort(
                "Falta motivo_seguimiento/nueva_informacion."
            )

        results: list[dict[str, Any]] = []

        # ------------------------------------------------------------------
        # 1. Resolver las 38 candidatas originales de CG/CJ.
        # ------------------------------------------------------------------
        source_date_counts = Counter()

        for row in candidates:
            excel_row = int(row["row"])
            column = row["column"].strip()
            ordinal = int(row["clause_ordinal"])
            clause = row["clause"]
            event_type = row["event_type"].strip()
            motive = row["motive"].strip()
            source_key_raw = row["source_key"]
            source_key = base.canonicalize_source_key(source_key_raw)

            item_base = {
                "row": excel_row,
                "column": column,
                "clause_ordinal": ordinal,
                "source_key": source_key_raw,
                "event_type": event_type,
                "motive": motive,
                "ambito": "colectivo",
                "clause": clause,
                "fuente": base.source_string(excel_row, column, ordinal),
                "origin": "CG_CJ",
            }

            pn = pn_index.get(source_key)
            if pn is None:
                results.append({
                    **item_base,
                    "resolution": "SKIP_SOURCE",
                    "id_proyecto_nucleo": None,
                    "target_type": "",
                    "target_id": None,
                    "fecha_evento": "",
                    "date_resolution": "UNRESOLVED",
                    "existing_ids": [],
                    "detail": "ProyectoNucleo no existe.",
                })
                continue

            full_dates = base.parse_full_dates(clause)
            resolved_date = None
            date_detail = ""

            if len(full_dates) == 1:
                resolved_date = full_dates[0]
                date_detail = (
                    "Fecha con año explícito en la cláusula: "
                    f"{resolved_date.isoformat()}."
                )
                source_date_counts["DATE_EXPLICIT"] += 1

            elif len(full_dates) > 1:
                results.append({
                    **item_base,
                    "resolution": "REVIEW_DATE",
                    "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                    "target_type": "",
                    "target_id": None,
                    "fecha_evento": "",
                    "date_resolution": "MULTIPLE_FULL_DATES",
                    "existing_ids": [],
                    "detail": (
                        f"La cláusula contiene {len(full_dates)} fechas "
                        "con año explícito."
                    ),
                })
                source_date_counts["REVIEW_DATE"] += 1
                continue

            elif event_type == "continuacion_asamblea":
                resolved_date, date_detail = corroborate_continuation_date_safe(
                    sheet,
                    excel_row,
                    clause,
                )
                if resolved_date is not None:
                    source_date_counts["DATE_CROSS_COLUMN"] += 1
                else:
                    results.append({
                        **item_base,
                        "resolution": "REVIEW_DATE",
                        "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                        "target_type": "",
                        "target_id": None,
                        "fecha_evento": "",
                        "date_resolution": "CROSS_COLUMN_UNRESOLVED",
                        "existing_ids": [],
                        "detail": date_detail,
                    })
                    source_date_counts["REVIEW_DATE"] += 1
                    continue

            else:
                results.append({
                    **item_base,
                    "resolution": "REVIEW_DATE",
                    "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                    "target_type": "",
                    "target_id": None,
                    "fecha_evento": "",
                    "date_resolution": "YEAR_MISSING",
                    "existing_ids": [],
                    "detail": (
                        "La cláusula describe un evento, pero no contiene "
                        "una fecha con año explícito. No se infiere el año."
                    ),
                })
                source_date_counts["REVIEW_DATE"] += 1
                continue

            if event_type == "continuacion_asamblea":
                assembly, group = resolve_2br_base_assembly(
                    db,
                    source_key=source_key,
                    pn=pn,
                    cop_catalog=cop_catalog,
                    context_catalog=context_catalog,
                    assembly_type_catalog=assembly_type_catalog,
                    result_catalog=result_catalog,
                )

                expected_date = group["continuation"]["date"]
                if resolved_date != expected_date:
                    raise ResolveAbort(
                        "Fecha de continuación CG/CJ no coincide con el plan 2B-R: "
                        f"{source_key}; resolver={resolved_date}, "
                        f"2B-R={expected_date}."
                    )

                target_type = "asamblea"
                target_id = assembly.id_asamblea
                target_detail = (
                    "Asamblea base resuelta por firma semántica congelada 2B-R; "
                    "la fecha de continuación pertenece al SeguimientoEvento, "
                    "no a una segunda Asamblea."
                )
            else:
                target_type = "proyecto_nucleo"
                target_id = pn.id_proyecto_nucleo
                target_detail = (
                    "Evento funcional de alcance colectivo vinculado "
                    "explícitamente al ProyectoNucleo."
                )

            event_type_id = event_catalog[event_type].id_catalogo_opcion
            motive_id = (
                motive_catalog[motive].id_catalogo_opcion
                if motive
                else None
            )

            existing = exact_existing_event(
                db,
                pn_id=pn.id_proyecto_nucleo,
                target_type=target_type,
                target_id=target_id,
                event_type_id=event_type_id,
                motive_id=motive_id,
                event_date=resolved_date,
                detail=clause,
                source=item_base["fuente"],
            )

            if len(existing) == 0:
                resolution = "CREATE_CANDIDATE"
                existing_ids = []
                existing_detail = "No existe SeguimientoEvento exacto."
            elif len(existing) == 1:
                resolution = "REUSE_CANDIDATE"
                existing_ids = [existing[0].id_seguimiento_evento]
                existing_detail = "SeguimientoEvento exacto ya existe."
            else:
                resolution = "REVIEW_EXISTING"
                existing_ids = [
                    item.id_seguimiento_evento for item in existing
                ]
                existing_detail = (
                    f"Hay {len(existing)} SeguimientoEvento exactos activos."
                )

            results.append({
                **item_base,
                "resolution": resolution,
                "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                "target_type": target_type,
                "target_id": target_id,
                "fecha_evento": resolved_date.isoformat(),
                "date_resolution": "DETERMINISTIC",
                "existing_ids": existing_ids,
                "detail": (
                    f"{date_detail} {target_detail} {existing_detail}"
                ),
            })

        review_date = [
            item
            for item in results
            if item["resolution"] == "REVIEW_DATE"
        ]
        deterministic_cgcj = [
            item
            for item in results
            if item["date_resolution"] == "DETERMINISTIC"
        ]

        if len(review_date) != EXPECTED_REVIEW_DATE:
            raise ResolveAbort(
                f"Cambió REVIEW_DATE CG/CJ: "
                f"esperado={EXPECTED_REVIEW_DATE}, actual={len(review_date)}."
            )
        if len(deterministic_cgcj) != EXPECTED_CG_CJ_DETERMINISTIC:
            raise ResolveAbort(
                "Cambió conjunto determinista CG/CJ: "
                f"esperado={EXPECTED_CG_CJ_DETERMINISTIC}, "
                f"actual={len(deterministic_cgcj)}."
            )

        # ------------------------------------------------------------------
        # 2. Añadir las cuatro continuaciones AI/AJ auditadas por 2B-R.
        # ------------------------------------------------------------------
        extras = []

        for group in repair2br.PLAN:
            continuation = group["continuation"]
            key = (
                group["source_key"],
                continuation["source_column"],
                continuation["source_row"],
            )

            # Las dos de CJ ya se resolvieron dentro de las 38 candidatas.
            if continuation["source_column"] == "CJ":
                continue

            if key not in EXPECTED_EXTRA_KEYS:
                raise ResolveAbort(
                    f"Continuación extra fuera del plan esperado: {key!r}."
                )

            pn = pn_index.get(group["source_key"])
            if pn is None:
                raise ResolveAbort(
                    f"ProyectoNucleo faltante para continuación {group['source_key']}."
                )

            assembly, resolved_group = resolve_2br_base_assembly(
                db,
                source_key=group["source_key"],
                pn=pn,
                cop_catalog=cop_catalog,
                context_catalog=context_catalog,
                assembly_type_catalog=assembly_type_catalog,
                result_catalog=result_catalog,
            )
            if resolved_group is not group:
                # Defensivo; ambos objetos deben provenir del mismo PLAN importado.
                raise ResolveAbort("Inconsistencia interna del plan 2B-R.")

            row_number = continuation["source_row"]
            column = continuation["source_column"]
            cell_value = sheet[f"{column}{row_number}"].value
            clause = continuation_fragment(cell_value)
            event_date = continuation["date"]
            source = extra_source_string(row_number, column)

            # La fecha debe estar explícita en la celda o corroborada exactamente
            # por AM de la misma fila; no se infiere el año libremente.
            full_dates = base.parse_full_dates(clause)
            if len(full_dates) == 1:
                if full_dates[0] != event_date:
                    raise ResolveAbort(
                        "Fecha explícita de continuación extra no coincide con 2B-R: "
                        f"{group['source_key']} {full_dates[0]} != {event_date}."
                    )
                date_resolution = "EXPLICIT_2BR_SOURCE"
                date_detail = (
                    f"Fecha explícita en {column}: {event_date.isoformat()}."
                )
            elif len(full_dates) == 0:
                partials = base.parse_partial_day_month(clause)
                am_dates = parse_excel_full_dates(sheet[f"AM{row_number}"].value)
                matches = [
                    value
                    for value in am_dates
                    if (value.day, value.month)
                    in set(partials)
                ]
                if len(partials) != 1 or matches != [event_date]:
                    raise ResolveAbort(
                        "La fecha parcial de continuación extra no se corroboró "
                        f"de forma única contra AM: {group['source_key']}."
                    )
                date_resolution = "CROSS_COLUMN_2BR_SOURCE"
                date_detail = (
                    f"Día/mes en {column} corroborado contra AM: "
                    f"{event_date.isoformat()}."
                )
            else:
                raise ResolveAbort(
                    "La continuación extra contiene múltiples fechas completas: "
                    f"{group['source_key']} / {column}{row_number}."
                )

            event_type_id = event_catalog[
                "continuacion_asamblea"
            ].id_catalogo_opcion

            existing = exact_existing_event(
                db,
                pn_id=pn.id_proyecto_nucleo,
                target_type="asamblea",
                target_id=assembly.id_asamblea,
                event_type_id=event_type_id,
                motive_id=None,
                event_date=event_date,
                detail=clause,
                source=source,
            )

            if len(existing) == 0:
                resolution = "CREATE_CANDIDATE"
                existing_ids = []
                existing_detail = "No existe SeguimientoEvento exacto."
            elif len(existing) == 1:
                resolution = "REUSE_CANDIDATE"
                existing_ids = [existing[0].id_seguimiento_evento]
                existing_detail = "SeguimientoEvento exacto ya existe."
            else:
                resolution = "REVIEW_EXISTING"
                existing_ids = [
                    item.id_seguimiento_evento for item in existing
                ]
                existing_detail = (
                    f"Hay {len(existing)} SeguimientoEvento exactos activos."
                )

            extra = {
                "row": row_number,
                "column": column,
                "clause_ordinal": 0,
                "source_key": group["source_key"],
                "event_type": "continuacion_asamblea",
                "motive": "",
                "ambito": "colectivo",
                "resolution": resolution,
                "id_proyecto_nucleo": pn.id_proyecto_nucleo,
                "target_type": "asamblea",
                "target_id": assembly.id_asamblea,
                "fecha_evento": event_date.isoformat(),
                "date_resolution": date_resolution,
                "existing_ids": existing_ids,
                "fuente": source,
                "clause": clause,
                "detail": (
                    f"{date_detail} Asamblea base resuelta por firma semántica "
                    f"congelada 2B-R ({continuation['assembly_label']}). "
                    f"{existing_detail}"
                ),
                "origin": "2B_R_AI_AJ",
            }
            extras.append(extra)

        observed_extra_keys = {
            (item["source_key"], item["column"], item["row"])
            for item in extras
        }
        if observed_extra_keys != EXPECTED_EXTRA_KEYS:
            raise ResolveAbort(
                "Cambió el conjunto de cuatro continuaciones extra 2B-R."
            )
        if len(extras) != EXPECTED_EXTRA_CONTINUATIONS:
            raise ResolveAbort(
                f"Continuaciones extra != {EXPECTED_EXTRA_CONTINUATIONS}."
            )

        results.extend(extras)

        deterministic_final = [
            item
            for item in results
            if item["date_resolution"] != "UNRESOLVED"
            and item["resolution"] != "REVIEW_DATE"
        ]

        if len(deterministic_final) != EXPECTED_FINAL_DETERMINISTIC:
            raise ResolveAbort(
                "Cambió deterministas finales 2J: "
                f"esperado={EXPECTED_FINAL_DETERMINISTIC}, "
                f"actual={len(deterministic_final)}."
            )

        final_by_event = Counter(
            (item["event_type"], item["motive"])
            for item in deterministic_final
        )
        for key, expected in EXPECTED_FINAL_BY_EVENT.items():
            if final_by_event[key] != expected:
                raise ResolveAbort(
                    f"Cambió determinista final {key}: "
                    f"esperado={expected}, actual={final_by_event[key]}."
                )
        extra_event_types = set(final_by_event) - set(EXPECTED_FINAL_BY_EVENT)
        if extra_event_types:
            raise ResolveAbort(
                "Aparecieron tipos deterministas fuera del plan final: "
                + ", ".join(map(str, sorted(extra_event_types)))
            )

        resolution_counts = Counter(item["resolution"] for item in results)

        # Ningún objetivo determinista debe seguir bloqueado o ambiguo.
        for code in (
            "BLOCKED_TARGET",
            "REVIEW_TARGET",
            "SKIP_SOURCE",
            "REVIEW_EXISTING",
        ):
            if resolution_counts[code] != 0:
                raise ResolveAbort(
                    f"2J post-reparaciones contiene {code}="
                    f"{resolution_counts[code]}."
                )

        create_count = resolution_counts["CREATE_CANDIDATE"]
        reuse_count = resolution_counts["REUSE_CANDIDATE"]

        if create_count + reuse_count != EXPECTED_FINAL_DETERMINISTIC:
            raise ResolveAbort(
                "CREATE+REUSE final 2J no coincide con 17: "
                f"CREATE={create_count}, REUSE={reuse_count}."
            )

        # Evitar aceptar un estado parcial silencioso.
        if (create_count, reuse_count) not in {
            (EXPECTED_FINAL_DETERMINISTIC, 0),
            (0, EXPECTED_FINAL_DETERMINISTIC),
        }:
            raise ResolveAbort(
                "Estado parcial 2J no permitido: "
                f"CREATE={create_count}, REUSE={reuse_count}."
            )

        fieldnames = [
            "row",
            "column",
            "clause_ordinal",
            "source_key",
            "event_type",
            "motive",
            "ambito",
            "resolution",
            "id_proyecto_nucleo",
            "target_type",
            "target_id",
            "fecha_evento",
            "date_resolution",
            "existing_ids",
            "fuente",
            "clause",
            "detail",
            "origin",
        ]

        with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=fieldnames)
            writer.writeheader()
            for item in results:
                writer.writerow({
                    **item,
                    "existing_ids": ";".join(
                        map(str, item["existing_ids"])
                    ),
                })

        summary = {
            "database": current_db,
            "schema": current_schema,
            "project": PROJECT_KEY,
            "excel": args.excel.name,
            "sha256": digest,
            "analysis_csv": args.analysis_csv.name,
            "cgcj": {
                "analysis_rows": len(analysis_rows),
                "event_candidates": len(candidates),
                "deterministic": len(deterministic_cgcj),
                "review_date": len(review_date),
                "date_explicit": source_date_counts["DATE_EXPLICIT"],
                "date_cross_column": source_date_counts["DATE_CROSS_COLUMN"],
            },
            "extra_continuations_2br": len(extras),
            "final": {
                "deterministic": len(deterministic_final),
                "CREATE_CANDIDATE": create_count,
                "REUSE_CANDIDATE": reuse_count,
                "REVIEW_DATE": resolution_counts["REVIEW_DATE"],
                "REVIEW_TARGET": 0,
                "BLOCKED_TARGET": 0,
                "REVIEW_EXISTING": 0,
                "SKIP_SOURCE": 0,
            },
            "deterministic_by_event": [
                {
                    "event_type": event_type,
                    "motive": motive,
                    "count": count,
                }
                for (event_type, motive), count
                in sorted(final_by_event.items())
            ],
            "continuations": [
                {
                    "source_key": item["source_key"],
                    "source": item["fuente"],
                    "date": item["fecha_evento"],
                    "target_type": item["target_type"],
                    "target_id": item["target_id"],
                    "resolution": item["resolution"],
                    "origin": item["origin"],
                }
                for item in results
                if item["event_type"] == "continuacion_asamblea"
                and item["resolution"] != "REVIEW_DATE"
            ],
        }

        json_path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        print("=== RESOLUCIÓN 2J POST-REPARACIONES — SEGUIMIENTO ===")
        print(f"BD: {current_db} / schema {current_schema}")
        print(f"Proyecto: {PROJECT_KEY}")
        print(f"SHA-256: {digest}")
        print("CG/CJ preservado:")
        print(f"  cláusulas análisis:          {len(analysis_rows)}")
        print(f"  EVENT_CANDIDATE:             {len(candidates)}")
        print(f"  deterministas CG/CJ:         {len(deterministic_cgcj)}")
        print(f"  REVIEW_DATE:                 {len(review_date)}")
        print("Continuaciones 2B-R:")
        print("  ya presentes en CJ:          2")
        print(f"  añadidas desde AI/AJ:        {len(extras)}")
        print("Plan final determinista:")
        print(f"  total:                       {len(deterministic_final)}")
        print(f"  CREATE_CANDIDATE:            {create_count}")
        print(f"  REUSE_CANDIDATE:             {reuse_count}")
        print(f"  BLOCKED_TARGET:              {resolution_counts['BLOCKED_TARGET']}")
        print(f"  REVIEW_TARGET:               {resolution_counts['REVIEW_TARGET']}")
        print(f"  REVIEW_EXISTING:             {resolution_counts['REVIEW_EXISTING']}")
        print(f"  SKIP_SOURCE:                 {resolution_counts['SKIP_SOURCE']}")
        print("Deterministas por tipo:")
        for (event_type, motive), count in sorted(final_by_event.items()):
            label = event_type + (f"/{motive}" if motive else "")
            print(f"  {label:<42} {count}")

        print("Continuaciones resueltas a Asamblea base:")
        for item in results:
            if (
                item["event_type"] == "continuacion_asamblea"
                and item["resolution"] != "REVIEW_DATE"
            ):
                print(
                    f"  {item['source_key']} | {item['fecha_evento']} | "
                    f"id_asamblea={item['target_id']} | {item['origin']}"
                )

        print(
            "Regla: la continuación no es una segunda Asamblea; "
            "la fecha pertenece al SeguimientoEvento."
        )
        print(
            "Regla: los 25 EVENT_CANDIDATE sin fecha determinista "
            "permanecen REVIEW_DATE."
        )
        print("No se creó SeguimientoEvento ni Documento.")
        print(f"Reporte JSON: {json_path}")
        print(f"Resolución CSV: {csv_path}")
        print("No se realizó ninguna escritura en PostgreSQL.")
        return 0

    finally:
        db.rollback()
        db.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (
        ResolveAbort,
        base.ResolveAbort,
        repair2br.ImportAbort,
        source2br.PreflightAbort,
    ) as exc:
        print("=== RESOLUCIÓN 2J POST-REPARACIONES ABORTADA ===", file=sys.stderr)
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
