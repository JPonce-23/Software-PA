#!/usr/bin/env python3
"""Importación transaccional 2D-R: reparar Convenio -> Asamblea autorizante.

Contexto:
- La capa 2D ya está cerrada e idempotente con 66 Convenios / 66 vínculos.
- La reparación 2B-R agregó 8 Asambleas y dejó cuatro Convenios previamente
  sin Asamblea con una asociación ahora inequívoca bajo la regla original 2D.
- Esta subcapa NO crea Convenios, NO crea Asambleas y NO modifica montos,
  superficies, fechas, afectaciones ni otros campos de negocio.
- Sólo completa id_asamblea_autorizacion para los cuatro casos congelados.

Regla preservada de 2D:
- id_asamblea_autorizacion sólo se usa cuando existe exactamente una Asamblea
  celebrada del mismo ProyectoNucleo + ciclo COP.
- La asociación se vuelve a resolver semánticamente; NO se congelan IDs de
  Asamblea obtenidos en ejecuciones previas.

Plan congelado posterior a 2B-R:
- Fuente 2D: 94 ciclos / 66 candidatos / 12 REVIEW / 2 SKIP / 14 NO_DATA.
- Asociación actual: UNIQUE=46 / NONE=16 / AMBIGUOUS=4.
- Estado persistido previo a 2D-R:
    42 Convenios ya correctamente enlazados.
    4 Convenios candidatos a UPDATE.
    16 siguen sin Asamblea.
    4 siguen ambiguos.
- Objetivos exactos:
    MEXICO|COYOTEPEC|COYOTEPEC / ORIGEN
    HIDALGO|CHAPANTONGO|JUCHITLAN / ORIGEN
    HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ / ORIGEN
    QUERETARO|SAN JUAN DEL RIO|PASO DE MATA / ORIGEN
- TULA DE ALLENDE permanece fuera por REVIEW/mixed_coverage.
- SANTIAGO OXTHOC permanece sin vínculo por AMBIGUOUS:2.

Seguridad:
- sólo DB_NAME=db_carga_excel
- schema 018
- SHA-256 exacto del Excel aprobado
- actor activo
- recompone y valida el plan fuente 2D
- exige exactamente 66 Convenios canónicos uno-a-uno
- sólo permite modificar los cuatro objetivos congelados
- valida que la Asamblea objetivo provenga de 2B-R y corresponda al grupo
- una sola transacción
- SET CONSTRAINTS ALL IMMEDIATE
- --dry-run => UPDATE real + validaciones + ROLLBACK
- --confirmar => COMMIT
- idempotencia: segunda ejecución => UPDATE=0 / REUSE=4
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.database import DB_NAME, SessionLocal  # noqa: E402
import importar_convenios_colectivos as base  # noqa: E402


EXPECTED_DB = "db_carga_excel"
EXPECTED_SCHEMA = "018"
EXPECTED_SHA256 = "bac01c25e9ff963f524fa8a48d69e74d92a0d36578731b1eca91f0e5a3573c3f"
PROJECT_KEY = "MEX-QRO"
SHEET = "INFORME M-Q"

EXPECTED_SOURCE = {
    "cycles": 94,
    "candidates": 66,
    "reviews": 12,
    "skips": 2,
    "no_data": 14,
}

EXPECTED_REVIEW_BREAKDOWN = {
    "mixed_coverage": 7,
    "affectation_ambiguous": 0,
    "date_text": 1,
    "multiple_sign_dates": 3,
    "multiple_program_dates": 0,
    "program_after_sign": 1,
    "bad_money": 0,
    "no_money": 0,
    "money_order": 0,
    "surface_mismatch": 0,
}

EXPECTED_ASSEMBLY_ASSOC_AFTER_2BR = {
    "unique": 46,
    "none": 16,
    "ambiguous": 4,
}

EXPECTED_UPDATE_KEYS = {
    ("MEXICO|COYOTEPEC|COYOTEPEC", "ORIGEN"),
    ("HIDALGO|CHAPANTONGO|JUCHITLAN", "ORIGEN"),
    ("HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ", "ORIGEN"),
    ("QUERETARO|SAN JUAN DEL RIO|PASO DE MATA", "ORIGEN"),
}

EXPECTED_TULA_KEY = ("HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE", "ORIGEN")
EXPECTED_SANTIAGO_KEY = ("MEXICO|JILOTEPEC|SANTIAGO OXTHOC", "ORIGEN")

EXPECTED_ALREADY_LINKED = 42
EXPECTED_UNLINKED_NONE = 16
EXPECTED_UNLINKED_AMBIGUOUS = 4
EXPECTED_TARGETS = 4


class ImportAbort(RuntimeError):
    pass


def exact_existing_convenios_ignoring_assembly(
    db,
    candidate: dict[str, Any],
):
    """Resuelve el Convenio 2D por su firma exacta, excepto Asamblea.

    Se conserva la relación ConvenioAfectacion principal como parte de la firma
    y no se toleran vínculos activos adicionales.
    """
    q = db.query(models.Convenio).filter(
        models.Convenio.id_proyecto_nucleo == candidate["pn_id"],
        models.Convenio.ambito == "colectivo",
        models.Convenio.tipo_instrumento == "convenio",
        models.Convenio.tipo_convenio == candidate["tipo_convenio"],
        models.Convenio.consecutivo == 1,
        models.Convenio.id_convenio_padre.is_(None),
        models.Convenio.efecto_monto == "pendiente",
        models.Convenio.monto_90_impacto.is_(None),
        models.Convenio.monto_100_impacto.is_(None),
        models.Convenio.monto_bdt_impacto.is_(None),
        models.Convenio.estado_antecedente == candidate["estado_antecedente"],
        models.Convenio.activo.is_(True),
    )

    fields = (
        ("fecha_programada_firma", models.Convenio.fecha_programada_firma),
        ("fecha_firma", models.Convenio.fecha_firma),
        ("monto_90", models.Convenio.monto_90),
        ("monto_100", models.Convenio.monto_100),
        ("monto_bdt", models.Convenio.monto_bdt),
        ("superficie_ha", models.Convenio.superficie_ha),
    )

    for name, column in fields:
        value = candidate[name]
        q = q.filter(column.is_(None)) if value is None else q.filter(column == value)

    matches = []
    for convenio in q.order_by(models.Convenio.id_convenio).all():
        links = (
            db.query(models.ConvenioAfectacion)
            .filter(
                models.ConvenioAfectacion.id_convenio == convenio.id_convenio,
                models.ConvenioAfectacion.activo.is_(True),
            )
            .order_by(models.ConvenioAfectacion.id_convenio_afectacion)
            .all()
        )
        if len(links) != 1:
            continue

        link = links[0]
        if (
            link.id_afectacion == candidate["affectation_id"]
            and link.rol == "principal"
            and link.efecto_superficie == "pendiente"
            and link.superficie_impacto_ha is None
        ):
            matches.append((convenio, link))

    return matches


def validate_source_plan(plan: dict[str, Any]) -> None:
    actual = {
        "cycles": len(plan["groups"]),
        "candidates": len(plan["candidates"]),
        "reviews": len(plan["reviews"]),
        "skips": len(plan["skips"]),
        "no_data": len(plan["no_data"]),
    }

    for key, expected in EXPECTED_SOURCE.items():
        if actual[key] != expected:
            raise ImportAbort(
                f"Cambió plan fuente 2D/{key}: esperado={expected}, actual={actual[key]}."
            )

    for key, expected in EXPECTED_REVIEW_BREAKDOWN.items():
        current = plan["review_stats"][key]
        if current != expected:
            raise ImportAbort(
                f"Cambió REVIEW 2D/{key}: esperado={expected}, actual={current}."
            )

    for key, expected in EXPECTED_ASSEMBLY_ASSOC_AFTER_2BR.items():
        current = plan["assembly_stats"][key]
        if current != expected:
            raise ImportAbort(
                "La asociación de Asambleas posterior a 2B-R no coincide: "
                f"{key} esperado={expected}, actual={current}."
            )


def active_project_convenio_count(db, project_id: int) -> int:
    return (
        db.query(models.Convenio)
        .join(
            models.ProyectoNucleo,
            models.ProyectoNucleo.id_proyecto_nucleo
            == models.Convenio.id_proyecto_nucleo,
        )
        .filter(
            models.ProyectoNucleo.id_proyecto == project_id,
            models.ProyectoNucleo.activo.is_(True),
            models.Convenio.activo.is_(True),
            models.Convenio.ambito == "colectivo",
            models.Convenio.tipo_instrumento == "convenio",
        )
        .count()
    )


def verify_repair_target(
    db,
    *,
    key: tuple[str, str],
    candidate: dict[str, Any],
) -> models.Asamblea:
    target_id = candidate["id_asamblea_autorizacion"]
    if target_id is None or candidate["assembly_status"] != "UNIQUE":
        raise ImportAbort(
            f"Objetivo 2D-R dejó de ser UNIQUE para {key[0]} / {key[1]}."
        )

    target = db.get(models.Asamblea, target_id)
    if target is None or not target.activo:
        raise ImportAbort(
            f"Asamblea objetivo inexistente/inactiva para {key[0]} / {key[1]}."
        )

    obs = target.observaciones or ""
    if (
        "Reparación migratoria 2B-R" not in obs
        or candidate["source_key"] not in obs
    ):
        raise ImportAbort(
            "La Asamblea única no puede acreditarse como objetivo creado por "
            f"2B-R para {key[0]} / {key[1]}."
        )

    if target.id_proyecto_nucleo != candidate["pn_id"]:
        raise ImportAbort(
            f"Asamblea objetivo no corresponde al ProyectoNucleo de {key[0]}."
        )

    return target


def classify_and_validate(
    db,
    plan: dict[str, Any],
) -> tuple[list[dict[str, Any]], Counter]:
    """Resuelve los 66 Convenios y congela las cuatro acciones reparables."""
    details: list[dict[str, Any]] = []
    counters = Counter()
    matched_convenio_ids: set[int] = set()
    observed_target_keys: set[tuple[str, str]] = set()

    for candidate in plan["candidates"]:
        key = (candidate["source_key"], candidate["cop_code"])
        matches = exact_existing_convenios_ignoring_assembly(db, candidate)

        if len(matches) != 1:
            raise ImportAbort(
                "La firma canónica de 2D sin Asamblea debe resolver exactamente "
                f"un Convenio para {key[0]} / {key[1]}; encontrados={len(matches)}."
            )

        convenio, _link = matches[0]
        if convenio.id_convenio in matched_convenio_ids:
            raise ImportAbort(
                f"El Convenio {convenio.id_convenio} resolvió a más de un candidato 2D."
            )
        matched_convenio_ids.add(convenio.id_convenio)

        current_id = convenio.id_asamblea_autorizacion
        target_id = candidate["id_asamblea_autorizacion"]
        assembly_status = candidate["assembly_status"]

        if key in EXPECTED_UPDATE_KEYS:
            observed_target_keys.add(key)
            verify_repair_target(db, key=key, candidate=candidate)

            if current_id is None:
                action = "UPDATE"
            elif current_id == target_id:
                action = "REUSE"
            else:
                raise ImportAbort(
                    f"Convenio {convenio.id_convenio} ya apunta a Asamblea "
                    f"{current_id}, pero 2D-R resuelve {target_id}."
                )

        elif target_id is not None:
            if assembly_status != "UNIQUE":
                raise ImportAbort(
                    f"Estado inconsistente UNIQUE/target para {key[0]} / {key[1]}."
                )
            if current_id != target_id:
                raise ImportAbort(
                    "Un Convenio fuera de la reparación focal ya no conserva su "
                    f"Asamblea 2D: {key[0]} / {key[1]} / "
                    f"actual={current_id} esperado={target_id}."
                )
            action = "ALREADY_LINKED"

        else:
            if current_id is not None:
                raise ImportAbort(
                    "Un Convenio sin asociación única tiene Asamblea persistida: "
                    f"{key[0]} / {key[1]} / id={current_id}."
                )
            if assembly_status == "NONE":
                action = "UNLINKED_NONE"
            elif assembly_status.startswith("AMBIGUOUS:"):
                action = "UNLINKED_AMBIGUOUS"
            else:
                raise ImportAbort(
                    f"Estado de Asamblea inesperado para {key[0]}: {assembly_status}."
                )

        counters[action] += 1
        details.append({
            "action": action,
            "source_key": candidate["source_key"],
            "cop": candidate["cop_code"],
            "source_rows": candidate["source_rows"],
            "id_convenio": convenio.id_convenio,
            "current_id_asamblea": current_id,
            "target_id_asamblea": target_id,
            "assembly_status": assembly_status,
        })

    if len(matched_convenio_ids) != 66:
        raise ImportAbort(
            "No se resolvieron uno-a-uno los 66 Convenios 2D: "
            f"resueltos={len(matched_convenio_ids)}."
        )

    if observed_target_keys != EXPECTED_UPDATE_KEYS:
        raise ImportAbort(
            "Cambió el conjunto exacto de objetivos 2D-R: "
            f"esperado={sorted(EXPECTED_UPDATE_KEYS)}, "
            f"actual={sorted(observed_target_keys)}."
        )

    if counters["ALREADY_LINKED"] != EXPECTED_ALREADY_LINKED:
        raise ImportAbort(
            "Cambió ALREADY_LINKED: "
            f"esperado={EXPECTED_ALREADY_LINKED}, "
            f"actual={counters['ALREADY_LINKED']}."
        )
    if counters["UNLINKED_NONE"] != EXPECTED_UNLINKED_NONE:
        raise ImportAbort(
            "Cambió UNLINKED_NONE: "
            f"esperado={EXPECTED_UNLINKED_NONE}, "
            f"actual={counters['UNLINKED_NONE']}."
        )
    if counters["UNLINKED_AMBIGUOUS"] != EXPECTED_UNLINKED_AMBIGUOUS:
        raise ImportAbort(
            "Cambió UNLINKED_AMBIGUOUS: "
            f"esperado={EXPECTED_UNLINKED_AMBIGUOUS}, "
            f"actual={counters['UNLINKED_AMBIGUOUS']}."
        )
    if counters["UPDATE"] + counters["REUSE"] != EXPECTED_TARGETS:
        raise ImportAbort(
            "UPDATE+REUSE de objetivos 2D-R debe ser exactamente 4."
        )

    santiago = [
        item for item in details
        if (item["source_key"], item["cop"]) == EXPECTED_SANTIAGO_KEY
    ]
    if (
        len(santiago) != 1
        or santiago[0]["action"] != "UNLINKED_AMBIGUOUS"
        or santiago[0]["assembly_status"] != "AMBIGUOUS:2"
    ):
        raise ImportAbort(
            "Santiago Oxthoc dejó de ser exactamente AMBIGUOUS:2."
        )

    return details, counters


def validate_tula_preserved(plan: dict[str, Any]) -> None:
    tula_candidate = any(
        (c["source_key"], c["cop_code"]) == EXPECTED_TULA_KEY
        for c in plan["candidates"]
    )
    tula_review = any(
        (r["source_key"], r["cop_code"]) == EXPECTED_TULA_KEY
        and r.get("reason_code") == "mixed_coverage"
        for r in plan["reviews"]
    )
    if tula_candidate or not tula_review:
        raise ImportAbort(
            "Tula de Allende cambió de clasificación 2D; "
            "se esperaba REVIEW/mixed_coverage."
        )


def postcheck_targets(
    db,
    plan: dict[str, Any],
) -> list[dict[str, Any]]:
    """Verifica que los cuatro Convenios apunten a la Asamblea semántica actual."""
    verified = []

    by_key = {
        (candidate["source_key"], candidate["cop_code"]): candidate
        for candidate in plan["candidates"]
    }

    for key in sorted(EXPECTED_UPDATE_KEYS):
        candidate = by_key.get(key)
        if candidate is None:
            raise ImportAbort(
                f"Postcheck: falta candidato 2D-R {key[0]} / {key[1]}."
            )

        verify_repair_target(db, key=key, candidate=candidate)

        matches = exact_existing_convenios_ignoring_assembly(db, candidate)
        if len(matches) != 1:
            raise ImportAbort(
                f"Postcheck: Convenio no resuelve único para {key[0]} / {key[1]}."
            )

        convenio, _link = matches[0]
        target_id = candidate["id_asamblea_autorizacion"]
        if convenio.id_asamblea_autorizacion != target_id:
            raise ImportAbort(
                f"Postcheck: Convenio {convenio.id_convenio} no quedó enlazado "
                f"a Asamblea {target_id}."
            )

        verified.append({
            "source_key": key[0],
            "cop": key[1],
            "id_convenio": convenio.id_convenio,
            "id_asamblea": target_id,
        })

    if len(verified) != EXPECTED_TARGETS:
        raise ImportAbort("Postcheck: no se verificaron los cuatro objetivos.")

    return verified


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)

    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--confirmar", action="store_true")

    parser.add_argument("--actor-id", type=int, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    args = parser.parse_args()

    if DB_NAME != EXPECTED_DB:
        raise ImportAbort(
            f"Protección: DB_NAME debe ser {EXPECTED_DB!r}; recibido {DB_NAME!r}."
        )

    if not args.excel.exists():
        raise ImportAbort(f"No existe el Excel: {args.excel}")

    digest = base.sha256_file(args.excel)
    if digest != EXPECTED_SHA256:
        raise ImportAbort(
            "El SHA-256 no coincide con el Excel aprobado. "
            f"Esperado {EXPECTED_SHA256}; recibido {digest}."
        )

    workbook = load_workbook(args.excel, data_only=True, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise ImportAbort(f"Falta la hoja {SHEET!r}.")
    sheet = workbook[SHEET]

    args.output_dir.mkdir(parents=True, exist_ok=True)
    mode_name = "confirmar" if args.confirmar else "dry-run"
    report_path = (
        args.output_dir
        / f"importacion_reparacion_convenios_asamblea_2d_r_{mode_name}.json"
    )

    db = SessionLocal()

    try:
        current_db = db.execute(text("SELECT current_database()")).scalar()
        current_schema = base.schema_version(db)

        if current_db != EXPECTED_DB or current_schema != EXPECTED_SCHEMA:
            raise ImportAbort(
                f"Destino inesperado: BD={current_db!r}, "
                f"schema={current_schema!r}."
            )

        actor = (
            db.query(models.Usuario)
            .filter(
                models.Usuario.id_usuario == args.actor_id,
                models.Usuario.activo.is_(True),
            )
            .one_or_none()
        )
        if actor is None:
            raise ImportAbort(
                f"Actor {args.actor_id} no existe o no está activo."
            )

        project = (
            db.query(models.Proyecto)
            .filter(
                models.Proyecto.clave_proyecto == PROJECT_KEY,
                models.Proyecto.activo.is_(True),
            )
            .one_or_none()
        )
        if project is None:
            raise ImportAbort(f"No existe proyecto activo {PROJECT_KEY!r}.")

        if active_project_convenio_count(db, project.id_proyecto) != 66:
            raise ImportAbort(
                "El laboratorio ya no contiene exactamente los 66 Convenios "
                "colectivos esperados de 2D."
            )

        pn_index = base.build_project_nucleus_index(db, project)
        cop_options = base.cop_catalog(db)
        celebrated_id = base.celebrated_result_id(db)

        for code in base.CONVENIO_TYPE_BY_COP:
            if code not in cop_options:
                raise ImportAbort(f"Falta tipo_cop_operativo/{code}.")

        plan = base.build_plan(
            db,
            sheet,
            pn_index,
            cop_options,
            celebrated_id,
        )

        validate_source_plan(plan)
        validate_tula_preserved(plan)

        before_details, before = classify_and_validate(db, plan)

        # Contexto de auditoría requerido por los triggers de bitácora.
        db.execute(
            text("SELECT set_config('app.current_user_id', :id, true)"),
            {"id": str(args.actor_id)},
        )

        updates = []
        now = datetime.now(timezone.utc)

        candidate_by_key = {
            (candidate["source_key"], candidate["cop_code"]): candidate
            for candidate in plan["candidates"]
        }

        for item in before_details:
            if item["action"] != "UPDATE":
                continue

            key = (item["source_key"], item["cop"])
            candidate = candidate_by_key[key]
            matches = exact_existing_convenios_ignoring_assembly(db, candidate)
            if len(matches) != 1:
                raise ImportAbort(
                    f"Objetivo dejó de resolver único durante UPDATE: {key}."
                )

            convenio, _link = matches[0]

            if convenio.id_asamblea_autorizacion is not None:
                raise ImportAbort(
                    f"Convenio {convenio.id_convenio} ya no tiene Asamblea NULL."
                )

            target_id = candidate["id_asamblea_autorizacion"]
            verify_repair_target(db, key=key, candidate=candidate)

            convenio.id_asamblea_autorizacion = target_id
            convenio.actualizado_en = now
            convenio.actualizado_por = args.actor_id
            db.flush()

            updates.append({
                "source_key": key[0],
                "cop": key[1],
                "id_convenio": convenio.id_convenio,
                "id_asamblea": target_id,
            })

        # Las restricciones diferidas deben evaluarse también durante dry-run.
        db.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))

        # Recalcular el plan después del UPDATE y verificar la persistencia
        # semántica de los cuatro enlaces, sin depender de IDs congelados.
        plan_after = base.build_plan(
            db,
            sheet,
            pn_index,
            cop_options,
            celebrated_id,
        )
        validate_source_plan(plan_after)
        validate_tula_preserved(plan_after)

        verified = postcheck_targets(db, plan_after)

        # Después de la reparación, una nueva clasificación debe ver los cuatro
        # objetivos como REUSE, incluso dentro del mismo dry-run antes del rollback.
        after_details, after = classify_and_validate(db, plan_after)
        if after["UPDATE"] != 0 or after["REUSE"] != EXPECTED_TARGETS:
            raise ImportAbort(
                "Postcheck de idempotencia intratransacción falló: "
                f"UPDATE={after['UPDATE']} REUSE={after['REUSE']}."
            )

        if len(updates) != before["UPDATE"]:
            raise ImportAbort(
                "El número real de UPDATE no coincide con el plan previo."
            )

        report = {
            "mode": mode_name,
            "database": current_db,
            "schema": current_schema,
            "actor": {
                "id_usuario": actor.id_usuario,
                "correo": actor.correo,
            },
            "project": PROJECT_KEY,
            "excel": args.excel.name,
            "sha256": digest,
            "source_plan": EXPECTED_SOURCE,
            "assembly_association": {
                key: plan_after["assembly_stats"][key]
                for key in ("unique", "none", "ambiguous")
            },
            "before": {
                "UPDATE": before["UPDATE"],
                "REUSE": before["REUSE"],
                "ALREADY_LINKED": before["ALREADY_LINKED"],
                "UNLINKED_NONE": before["UNLINKED_NONE"],
                "UNLINKED_AMBIGUOUS": before["UNLINKED_AMBIGUOUS"],
            },
            "updates": updates,
            "verified": verified,
            "after_in_transaction": {
                "UPDATE": after["UPDATE"],
                "REUSE": after["REUSE"],
                "ALREADY_LINKED": after["ALREADY_LINKED"],
                "UNLINKED_NONE": after["UNLINKED_NONE"],
                "UNLINKED_AMBIGUOUS": after["UNLINKED_AMBIGUOUS"],
            },
            "tula_review_preserved": True,
            "santiago_ambiguous_preserved": True,
        }

        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

        print("=== IMPORTACIÓN REPARACIÓN 2D-R — CONVENIO -> ASAMBLEA ===")
        print(f"Modo: {mode_name}")
        print(f"BD: {current_db} / schema {current_schema}")
        print(f"Actor: {actor.id_usuario} / {actor.correo}")
        print(f"Proyecto: {PROJECT_KEY}")
        print(f"SHA-256: {digest}")
        print("Asociación de Asambleas preservada:")
        print(f"  UNIQUE:       {plan_after['assembly_stats']['unique']}")
        print(f"  NONE:         {plan_after['assembly_stats']['none']}")
        print(f"  AMBIGUOUS:    {plan_after['assembly_stats']['ambiguous']}")
        print("Objetivos 2D-R al inicio:")
        print(f"  UPDATE   {before['UPDATE']}")
        print(f"  REUSE    {before['REUSE']}")
        print("Casos no modificados:")
        print(f"  ALREADY_LINKED      {before['ALREADY_LINKED']}")
        print(f"  UNLINKED_NONE       {before['UNLINKED_NONE']}")
        print(f"  UNLINKED_AMBIGUOUS  {before['UNLINKED_AMBIGUOUS']}")
        print("Validación intratransacción:")
        print(f"  UPDATE   {after['UPDATE']}")
        print(f"  REUSE    {after['REUSE']}")
        print("Enlaces 2D-R verificados:")
        for item in verified:
            print(
                f"  {item['source_key']} | {item['cop']} | "
                f"id_convenio={item['id_convenio']} -> "
                f"id_asamblea={item['id_asamblea']}"
            )
        print("Conservadores:")
        print("  TULA DE ALLENDE: REVIEW/mixed_coverage preservado")
        print("  SANTIAGO OXTHOC: AMBIGUOUS:2 preservado")

        if args.confirmar:
            db.commit()
            print("COMMIT realizado")
        else:
            db.rollback()
            print("ROLLBACK realizado (dry-run)")

        print(f"Reporte: {report_path}")
        return 0

    except IntegrityError as exc:
        db.rollback()
        print("=== IMPORTACIÓN 2D-R ABORTADA ===", file=sys.stderr)
        print(f"IntegrityError: {exc}", file=sys.stderr)
        return 2
    except (ImportAbort, base.ImportAbort) as exc:
        db.rollback()
        print("=== IMPORTACIÓN 2D-R ABORTADA ===", file=sys.stderr)
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
