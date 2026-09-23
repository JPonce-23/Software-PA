#!/usr/bin/env python3
"""Preflight read-only 2D-R: reparar Asamblea autorizante en convenios colectivos.

Contexto:
- La capa 2D ya importó 66 Convenios / 66 ConvenioAfectacion.
- En el cierre original, la asociación a Asamblea quedó:
    UNIQUE=42, NONE=21, AMBIGUOUS=3.
- La reparación 2B-R agregó únicamente las Asambleas permanentes faltantes.
- Esta subcapa NO reimporta convenios ni altera montos/superficies/fechas.
- Sólo identifica Convenios existentes cuyo id_asamblea_autorizacion era NULL y
  que, después de 2B-R, ahora tienen exactamente una Asamblea celebrada del
  mismo ProyectoNucleo + ciclo COP según la regla original de 2D.

Plan esperado después del COMMIT idempotente de 2B-R:
- 66 candidatos canónicos de 2D siguen existiendo.
- Asociación recalculada contra las Asambleas actuales:
    UNIQUE=46, NONE=16, AMBIGUOUS=4.
- Estado de los Convenios ya persistidos:
    ALREADY_LINKED=42
    UPDATE_CANDIDATE=4
    UNLINKED_NONE=16
    UNLINKED_AMBIGUOUS=4
    REVIEW=0
- Los cuatro UPDATE_CANDIDATE esperados son:
    MEXICO|COYOTEPEC|COYOTEPEC / ORIGEN
    HIDALGO|CHAPANTONGO|JUCHITLAN / ORIGEN
    HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ / ORIGEN
    QUERETARO|SAN JUAN DEL RIO|PASO DE MATA / ORIGEN
- TULA DE ALLENDE no entra: 2D la dejó REVIEW por mixed_coverage.
- SANTIAGO OXTHOC permanece sin vínculo: ahora tiene 2 Asambleas celebradas
  para el mismo ProyectoNucleo + COP y por tanto continúa AMBIGUOUS.

Este script NO escribe en PostgreSQL.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

# /app/scripts -> /app para importar app.*; además permite importar el importador
# 2D ya cerrado y reutilizar exactamente su parser/reglas de fuente.
sys.path.append(str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.database import DB_NAME, SessionLocal  # noqa: E402
import importar_convenios_colectivos as base  # noqa: E402


EXPECTED_DB = "db_carga_excel"
EXPECTED_SCHEMA = "018"
EXPECTED_SHA256 = "bac01c25e9ff963f524fa8a48d69e74d92a0d36578731b1eca91f0e5a3573c3f"
PROJECT_KEY = "MEX-QRO"
SHEET = "INFORME M-Q"

EXPECTED_SOURCE = {
    "cycles": 94,
    "candidates": 66,
    "reviews": 12,
    "skips": 2,
    "no_data": 14,
}
EXPECTED_REVIEW_BREAKDOWN = {
    "mixed_coverage": 7,
    "affectation_ambiguous": 0,
    "date_text": 1,
    "multiple_sign_dates": 3,
    "multiple_program_dates": 0,
    "program_after_sign": 1,
    "bad_money": 0,
    "no_money": 0,
    "money_order": 0,
    "surface_mismatch": 0,
}
EXPECTED_ASSEMBLY_ASSOC_AFTER_2BR = {
    "unique": 46,
    "none": 16,
    "ambiguous": 4,
}
EXPECTED_ACTIONS = {
    "ALREADY_LINKED": 42,
    "UPDATE_CANDIDATE": 4,
    "UNLINKED_NONE": 16,
    "UNLINKED_AMBIGUOUS": 4,
    "REVIEW": 0,
}
EXPECTED_UPDATE_KEYS = {
    ("MEXICO|COYOTEPEC|COYOTEPEC", "ORIGEN"),
    ("HIDALGO|CHAPANTONGO|JUCHITLAN", "ORIGEN"),
    ("HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ", "ORIGEN"),
    ("QUERETARO|SAN JUAN DEL RIO|PASO DE MATA", "ORIGEN"),
}
EXPECTED_SANTIAGO_KEY = ("MEXICO|JILOTEPEC|SANTIAGO OXTHOC", "ORIGEN")
EXPECTED_TULA_KEY = ("HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE", "ORIGEN")


class PreflightAbort(RuntimeError):
    pass


def exact_existing_convenios_ignoring_assembly(
    db,
    candidate: dict[str, Any],
):
    """Encuentra el Convenio 2D por la misma firma canónica salvo Asamblea.

    La relación ConvenioAfectacion principal sí forma parte de la firma. No se
    toleran vínculos adicionales activos.
    """
    q = db.query(models.Convenio).filter(
        models.Convenio.id_proyecto_nucleo == candidate["pn_id"],
        models.Convenio.ambito == "colectivo",
        models.Convenio.tipo_instrumento == "convenio",
        models.Convenio.tipo_convenio == candidate["tipo_convenio"],
        models.Convenio.consecutivo == 1,
        models.Convenio.id_convenio_padre.is_(None),
        models.Convenio.efecto_monto == "pendiente",
        models.Convenio.monto_90_impacto.is_(None),
        models.Convenio.monto_100_impacto.is_(None),
        models.Convenio.monto_bdt_impacto.is_(None),
        models.Convenio.estado_antecedente == candidate["estado_antecedente"],
        models.Convenio.activo.is_(True),
    )

    fields = (
        ("fecha_programada_firma", models.Convenio.fecha_programada_firma),
        ("fecha_firma", models.Convenio.fecha_firma),
        ("monto_90", models.Convenio.monto_90),
        ("monto_100", models.Convenio.monto_100),
        ("monto_bdt", models.Convenio.monto_bdt),
        ("superficie_ha", models.Convenio.superficie_ha),
    )

    for name, column in fields:
        value = candidate[name]
        q = q.filter(column.is_(None)) if value is None else q.filter(column == value)

    matches = []
    for convenio in q.order_by(models.Convenio.id_convenio).all():
        links = (
            db.query(models.ConvenioAfectacion)
            .filter(
                models.ConvenioAfectacion.id_convenio == convenio.id_convenio,
                models.ConvenioAfectacion.activo.is_(True),
            )
            .order_by(models.ConvenioAfectacion.id_convenio_afectacion)
            .all()
        )
        if len(links) != 1:
            continue
        link = links[0]
        if (
            link.id_afectacion == candidate["affectation_id"]
            and link.rol == "principal"
            and link.efecto_superficie == "pendiente"
            and link.superficie_impacto_ha is None
        ):
            matches.append((convenio, link))

    return matches


def validate_source_plan(plan: dict[str, Any]) -> None:
    actual = {
        "cycles": len(plan["groups"]),
        "candidates": len(plan["candidates"]),
        "reviews": len(plan["reviews"]),
        "skips": len(plan["skips"]),
        "no_data": len(plan["no_data"]),
    }
    for key, expected in EXPECTED_SOURCE.items():
        if actual[key] != expected:
            raise PreflightAbort(
                f"Cambió plan fuente 2D/{key}: esperado={expected}, actual={actual[key]}."
            )

    for key, expected in EXPECTED_REVIEW_BREAKDOWN.items():
        current = plan["review_stats"][key]
        if current != expected:
            raise PreflightAbort(
                f"Cambió REVIEW 2D/{key}: esperado={expected}, actual={current}."
            )

    for key, expected in EXPECTED_ASSEMBLY_ASSOC_AFTER_2BR.items():
        current = plan["assembly_stats"][key]
        if current != expected:
            raise PreflightAbort(
                "La asociación de Asambleas posterior a 2B-R no coincide: "
                f"{key} esperado={expected}, actual={current}."
            )


def active_project_convenio_count(db, project_id: int) -> int:
    return (
        db.query(models.Convenio)
        .join(
            models.ProyectoNucleo,
            models.ProyectoNucleo.id_proyecto_nucleo
            == models.Convenio.id_proyecto_nucleo,
        )
        .filter(
            models.ProyectoNucleo.id_proyecto == project_id,
            models.ProyectoNucleo.activo.is_(True),
            models.Convenio.activo.is_(True),
            models.Convenio.ambito == "colectivo",
            models.Convenio.tipo_instrumento == "convenio",
        )
        .count()
    )


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    args = parser.parse_args()

    if DB_NAME != EXPECTED_DB:
        raise PreflightAbort(
            f"Protección: DB_NAME debe ser {EXPECTED_DB!r}; recibido {DB_NAME!r}."
        )
    if not args.excel.exists():
        raise PreflightAbort(f"No existe el Excel: {args.excel}")

    digest = base.sha256_file(args.excel)
    if digest != EXPECTED_SHA256:
        raise PreflightAbort(
            "El SHA-256 no coincide con el Excel aprobado. "
            f"Esperado {EXPECTED_SHA256}; recibido {digest}."
        )

    workbook = load_workbook(args.excel, data_only=True, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise PreflightAbort(f"Falta la hoja {SHEET!r}.")
    sheet = workbook[SHEET]

    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.output_dir / "preflight_reparacion_convenios_asamblea_2d_r.csv"
    json_path = args.output_dir / "preflight_reparacion_convenios_asamblea_2d_r.json"

    db = SessionLocal()
    try:
        current_db = db.execute(text("SELECT current_database()" )).scalar()
        current_schema = base.schema_version(db)
        if current_db != EXPECTED_DB or current_schema != EXPECTED_SCHEMA:
            raise PreflightAbort(
                f"Destino inesperado: BD={current_db!r}, schema={current_schema!r}."
            )

        project = (
            db.query(models.Proyecto)
            .filter(
                models.Proyecto.clave_proyecto == PROJECT_KEY,
                models.Proyecto.activo.is_(True),
            )
            .one_or_none()
        )
        if project is None:
            raise PreflightAbort(f"No existe proyecto activo {PROJECT_KEY!r}.")

        pn_index = base.build_project_nucleus_index(db, project)
        cop_options = base.cop_catalog(db)
        celebrated_id = base.celebrated_result_id(db)

        for code in base.CONVENIO_TYPE_BY_COP:
            if code not in cop_options:
                raise PreflightAbort(f"Falta tipo_cop_operativo/{code}.")

        plan = base.build_plan(
            db,
            sheet,
            pn_index,
            cop_options,
            celebrated_id,
        )
        validate_source_plan(plan)

        total_project_convenios = active_project_convenio_count(
            db, project.id_proyecto
        )
        if total_project_convenios != 66:
            raise PreflightAbort(
                "El laboratorio ya no contiene exactamente los 66 Convenios "
                f"colectivos activos de 2D: actual={total_project_convenios}."
            )

        actions = Counter()
        details = []
        matched_convenio_ids: set[int] = set()

        for candidate in plan["candidates"]:
            matches = exact_existing_convenios_ignoring_assembly(db, candidate)
            key = (candidate["source_key"], candidate["cop_code"])

            if len(matches) != 1:
                actions["REVIEW"] += 1
                details.append({
                    "action": "REVIEW",
                    "source_key": candidate["source_key"],
                    "cop": candidate["cop_code"],
                    "source_rows": candidate["source_rows"],
                    "id_convenio": None,
                    "current_id_asamblea": None,
                    "target_id_asamblea": candidate["id_asamblea_autorizacion"],
                    "assembly_status": candidate["assembly_status"],
                    "detail": (
                        "La firma canónica de 2D sin Asamblea resolvió "
                        f"{len(matches)} Convenios; se requiere revisión."
                    ),
                })
                continue

            convenio, _link = matches[0]
            if convenio.id_convenio in matched_convenio_ids:
                raise PreflightAbort(
                    f"El Convenio {convenio.id_convenio} resolvió a más de un candidato 2D."
                )
            matched_convenio_ids.add(convenio.id_convenio)

            target_id = candidate["id_asamblea_autorizacion"]
            current_id = convenio.id_asamblea_autorizacion
            assembly_status = candidate["assembly_status"]

            if target_id is not None:
                if current_id == target_id:
                    action = "ALREADY_LINKED"
                    detail = "Convenio ya conserva la Asamblea única esperada por 2D."
                elif current_id is None:
                    target = db.get(models.Asamblea, target_id)
                    obs = (target.observaciones or "") if target else ""
                    if key not in EXPECTED_UPDATE_KEYS:
                        action = "REVIEW"
                        detail = (
                            "Apareció una nueva asociación única fuera del conjunto "
                            "focal esperado de 2B-R."
                        )
                    elif (
                        target is None
                        or "Reparación migratoria 2B-R" not in obs
                        or candidate["source_key"] not in obs
                    ):
                        action = "REVIEW"
                        detail = (
                            "La Asamblea única no puede acreditarse como objetivo "
                            "creado por la reparación 2B-R para este grupo."
                        )
                    else:
                        action = "UPDATE_CANDIDATE"
                        detail = (
                            "Convenio existente con Asamblea NULL; 2B-R dejó una "
                            "única Asamblea celebrada determinista para el mismo PN+COP."
                        )
                else:
                    action = "REVIEW"
                    detail = (
                        f"Convenio ya apunta a id_asamblea={current_id}, pero la regla "
                        f"2D resuelve id_asamblea={target_id}."
                    )
            else:
                if current_id is not None:
                    action = "REVIEW"
                    detail = (
                        "Convenio tiene Asamblea persistida aunque la regla 2D actual "
                        f"no es única ({assembly_status})."
                    )
                elif assembly_status == "NONE":
                    action = "UNLINKED_NONE"
                    detail = "Sigue sin Asamblea celebrada única; no se modifica."
                elif assembly_status.startswith("AMBIGUOUS:"):
                    action = "UNLINKED_AMBIGUOUS"
                    detail = "Sigue/queda ambiguo; no se modifica."
                else:
                    action = "REVIEW"
                    detail = f"Estado de Asamblea inesperado: {assembly_status}."

            actions[action] += 1
            details.append({
                "action": action,
                "source_key": candidate["source_key"],
                "cop": candidate["cop_code"],
                "source_rows": candidate["source_rows"],
                "id_convenio": convenio.id_convenio,
                "current_id_asamblea": current_id,
                "target_id_asamblea": target_id,
                "assembly_status": assembly_status,
                "fecha_firma": (
                    candidate["fecha_firma"].isoformat()
                    if candidate["fecha_firma"] else None
                ),
                "detail": detail,
            })

        if len(matched_convenio_ids) != 66:
            raise PreflightAbort(
                "No se resolvieron de forma uno-a-uno los 66 Convenios de 2D: "
                f"resueltos={len(matched_convenio_ids)}."
            )

        for action, expected in EXPECTED_ACTIONS.items():
            actual = actions[action]
            if actual != expected:
                raise PreflightAbort(
                    f"Cambió resultado {action}: esperado={expected}, actual={actual}."
                )

        update_keys = {
            (item["source_key"], item["cop"])
            for item in details
            if item["action"] == "UPDATE_CANDIDATE"
        }
        if update_keys != EXPECTED_UPDATE_KEYS:
            raise PreflightAbort(
                "Cambió el conjunto exacto a reparar. "
                f"Esperado={sorted(EXPECTED_UPDATE_KEYS)}; actual={sorted(update_keys)}."
            )

        santiago = [
            item for item in details
            if (item["source_key"], item["cop"]) == EXPECTED_SANTIAGO_KEY
        ]
        if len(santiago) != 1 or santiago[0]["action"] != "UNLINKED_AMBIGUOUS":
            raise PreflightAbort(
                "Santiago Oxthoc no quedó como único caso focal AMBIGUOUS esperado."
            )
        if santiago[0]["assembly_status"] != "AMBIGUOUS:2":
            raise PreflightAbort(
                "Santiago Oxthoc debe resolver exactamente 2 Asambleas celebradas."
            )

        # Tula debe seguir fuera de los candidatos 2D por mixed_coverage.
        tula_candidate = any(
            (c["source_key"], c["cop_code"]) == EXPECTED_TULA_KEY
            for c in plan["candidates"]
        )
        tula_review = any(
            (r["source_key"], r["cop_code"]) == EXPECTED_TULA_KEY
            and r.get("reason_code") == "mixed_coverage"
            for r in plan["reviews"]
        )
        if tula_candidate or not tula_review:
            raise PreflightAbort(
                "Tula de Allende cambió de clasificación 2D; se esperaba REVIEW/mixed_coverage."
            )

        fieldnames = [
            "action",
            "source_key",
            "cop",
            "source_rows",
            "id_convenio",
            "current_id_asamblea",
            "target_id_asamblea",
            "assembly_status",
            "fecha_firma",
            "detail",
        ]
        with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=fieldnames)
            writer.writeheader()
            for item in details:
                row = dict(item)
                row["source_rows"] = ";".join(map(str, row["source_rows"]))
                writer.writerow(row)

        report = {
            "database": current_db,
            "schema": current_schema,
            "project": PROJECT_KEY,
            "sha256": digest,
            "source_plan": EXPECTED_SOURCE,
            "assembly_association_after_2br": {
                key: plan["assembly_stats"][key]
                for key in ("unique", "none", "ambiguous")
            },
            "actions": {key: actions[key] for key in EXPECTED_ACTIONS},
            "update_candidates": [
                item for item in details if item["action"] == "UPDATE_CANDIDATE"
            ],
            "unlinked_ambiguous": [
                item for item in details if item["action"] == "UNLINKED_AMBIGUOUS"
            ],
            "tula_review_preserved": True,
            "details": details,
            "writes": 0,
        }
        with json_path.open("w", encoding="utf-8") as stream:
            json.dump(report, stream, ensure_ascii=False, indent=2, default=str)

        print("=== PREFLIGHT REPARACIÓN 2D-R — CONVENIO -> ASAMBLEA ===")
        print(f"BD: {current_db} / schema {current_schema}")
        print(f"Proyecto: {PROJECT_KEY}")
        print(f"SHA-256: {digest}")
        print("Plan fuente 2D preservado:")
        print(f"  Ciclos:       {len(plan['groups'])}")
        print(f"  Candidatos:   {len(plan['candidates'])}")
        print(f"  REVIEW:       {len(plan['reviews'])}")
        print(f"  SKIP:         {len(plan['skips'])}")
        print(f"  NO_DATA:      {len(plan['no_data'])}")
        print("Asociación recalculada después de 2B-R:")
        print(f"  UNIQUE:       {plan['assembly_stats']['unique']}")
        print(f"  NONE:         {plan['assembly_stats']['none']}")
        print(f"  AMBIGUOUS:    {plan['assembly_stats']['ambiguous']}")
        print("Convenios persistidos:")
        for action in (
            "ALREADY_LINKED",
            "UPDATE_CANDIDATE",
            "UNLINKED_NONE",
            "UNLINKED_AMBIGUOUS",
            "REVIEW",
        ):
            print(f"  {action:<19} {actions[action]}")
        print("Candidatos exactos a reparación:")
        for item in details:
            if item["action"] == "UPDATE_CANDIDATE":
                print(
                    "  "
                    f"{item['source_key']} | {item['cop']} | "
                    f"id_convenio={item['id_convenio']} -> "
                    f"id_asamblea={item['target_id_asamblea']}"
                )
        print("Casos focales conservadores:")
        print("  TULA DE ALLENDE: REVIEW/mixed_coverage (sin Convenio 2D a reparar)")
        print("  SANTIAGO OXTHOC: AMBIGUOUS:2 (Convenio permanece sin Asamblea)")
        print(f"Reporte JSON: {json_path}")
        print(f"Plan CSV:     {csv_path}")
        print("No se realizó ninguna escritura en PostgreSQL.")

        return 0

    finally:
        db.rollback()
        db.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except PreflightAbort as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
    except base.ImportAbort as exc:
        print(f"ERROR BASE 2D: {exc}", file=sys.stderr)
        raise SystemExit(2)
