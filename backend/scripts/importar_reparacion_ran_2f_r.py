#!/usr/bin/env python3
"""Importación transaccional 2F-R: trámites RAN desbloqueados por 2B-R.

Esta reparación es focalizada. No reimporta ni modifica los 123 TramiteRan
históricos de 2F; únicamente crea, cuando todavía no existen, los ocho trámites
de asamblea de anuencia que pasaron de BLOCKED a deterministas después de 2B-R.

Plan congelado por preflight_reparacion_ran_2f_r.py:
- Fuente: 88 asamblea_anuencia / 88 convenio / 3 asamblea_retiro.
- Persistencia previa: 123 TramiteRan / 208 TramiteRanEvento.
- Nuevos deterministas: 8 TramiteRan de asamblea_anuencia.
- Eventos nuevos: 8 ingreso / 0 calificacion / 7 inscripcion = 15.
- Permanecen: REVIEW=6 / BLOCKED=15 / SKIP=2.
- Después de importar: 131 TramiteRan / 223 TramiteRanEvento.
- Segunda ejecución: CREATE=0 / REUSE=8 para los objetivos 2F-R.

Objetivos exactos 2F-R:
1. COYOTEPEC / ORIGEN / filas 45,47..59 / MATCH_REALIZED_DATE
2. TULA DE ALLENDE / ORIGEN / fila 92 / UNIQUE_CYCLE
3. TULA DE ALLENDE / ORIGEN / fila 93 / UNIQUE_CYCLE
4. JUCHITLAN / ORIGEN / fila 103 / UNIQUE_CYCLE
5. SAN JOSE EL MARQUEZ / ORIGEN / filas 105,106 / UNIQUE_CYCLE
6. SANTIAGO OXTHOC / ORIGEN / fila 117 / MATCH_REALIZED_DATE
7. PASO DE MATA / ORIGEN / fila 136 / UNIQUE_CYCLE
8. PASO DE MATA / ORIGEN / filas 138,139 / UNIQUE_CYCLE

Reglas preservadas:
- Se reutiliza exactamente el parser, deduplicación, resolución de objetivo y
  firma de idempotencia de importar_ran_colectivos.py.
- No se congelan IDs de Asamblea.
- Un objetivo nuevo sólo es válido si la Asamblea se acredita como creada por
  la reparación 2B-R para el mismo source_key.
- Los 14 BLOCKED de convenio deben seguir BLOCKED.
- LA CUEVA / ADICIONAL debe seguir BLOCKED en asamblea_anuencia.
- Los 6 REVIEW y 2 SKIP originales deben permanecer.
- Tula filas 92 y 93 son dos trámites distintos por firma de eventos.
- Santiago fila 117 corresponde a la Asamblea celebrada 2025-11-18.

Seguridad:
- sólo DB_NAME=db_carga_excel
- schema 018
- SHA-256 exacto
- actor activo
- proyecto MEX-QRO
- estado previo permitido: 123/208 (antes de reparar) o 131/223 (ya reparado)
- sólo se crean los ocho objetivos congelados
- una sola transacción
- SET CONSTRAINTS ALL IMMEDIATE
- --dry-run => INSERT real + validaciones + ROLLBACK
- --confirmar => COMMIT
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.database import DB_NAME, SessionLocal  # noqa: E402
import importar_ran_colectivos as base  # noqa: E402


EXPECTED_DB = "db_carga_excel"
EXPECTED_SCHEMA = "018"
EXPECTED_SHA256 = "bac01c25e9ff963f524fa8a48d69e74d92a0d36578731b1eca91f0e5a3573c3f"
PROJECT_KEY = "MEX-QRO"
SHEET = "INFORME M-Q"

EXPECTED_SOURCE_ROWS = {
    "asamblea_anuencia": 88,
    "convenio": 88,
    "asamblea_retiro": 3,
}

EXPECTED_REVIEW = 6
EXPECTED_BLOCKED = 15
EXPECTED_SKIP = 2

EXPECTED_FINAL_TRAMITES = 131
EXPECTED_FINAL_EVENTOS = 223

EXPECTED_FINAL_BY_OBJECTIVE = {
    "asamblea_anuencia": {
        "reuse": 55,
        "review": 6,
        "blocked": 1,
        "skip": 1,
    },
    "convenio": {
        "reuse": 73,
        "review": 0,
        "blocked": 14,
        "skip": 1,
    },
    "asamblea_retiro": {
        "reuse": 3,
        "review": 0,
        "blocked": 0,
        "skip": 0,
    },
}

EXPECTED_FINAL_EVENTS = {
    "ingreso": 131,
    "calificacion": 9,
    "inscripcion": 83,
}

EXPECTED_NEW_EVENTS = {
    "ingreso": 8,
    "calificacion": 0,
    "inscripcion": 7,
}

EXPECTED_TARGETS = {
    (
        "MEXICO|COYOTEPEC|COYOTEPEC",
        "ORIGEN",
        (45, 47, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59),
        "MATCH_REALIZED_DATE",
    ),
    (
        "HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE",
        "ORIGEN",
        (92,),
        "UNIQUE_CYCLE",
    ),
    (
        "HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE",
        "ORIGEN",
        (93,),
        "UNIQUE_CYCLE",
    ),
    (
        "HIDALGO|CHAPANTONGO|JUCHITLAN",
        "ORIGEN",
        (103,),
        "UNIQUE_CYCLE",
    ),
    (
        "HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ",
        "ORIGEN",
        (105, 106),
        "UNIQUE_CYCLE",
    ),
    (
        "MEXICO|JILOTEPEC|SANTIAGO OXTHOC",
        "ORIGEN",
        (117,),
        "MATCH_REALIZED_DATE",
    ),
    (
        "QUERETARO|SAN JUAN DEL RIO|PASO DE MATA",
        "ORIGEN",
        (136,),
        "UNIQUE_CYCLE",
    ),
    (
        "QUERETARO|SAN JUAN DEL RIO|PASO DE MATA",
        "ORIGEN",
        (138, 139),
        "UNIQUE_CYCLE",
    ),
}

EXPECTED_REMAINING_ANUENCIA_BLOCKED = {
    (
        "QUERETARO|SAN JUAN DEL RIO|LA CUEVA",
        "ADICIONAL",
        (133,),
        "NONE",
    )
}


class ImportAbort(RuntimeError):
    pass


def project_ran_counts(db, project_id: int) -> tuple[int, int]:
    tramites = (
        db.query(models.TramiteRan)
        .join(
            models.ProyectoNucleo,
            models.ProyectoNucleo.id_proyecto_nucleo
            == models.TramiteRan.id_proyecto_nucleo,
        )
        .filter(
            models.ProyectoNucleo.id_proyecto == project_id,
            models.ProyectoNucleo.activo.is_(True),
            models.TramiteRan.activo.is_(True),
        )
        .count()
    )

    eventos = (
        db.query(models.TramiteRanEvento)
        .join(
            models.TramiteRan,
            models.TramiteRan.id_tramite_ran
            == models.TramiteRanEvento.id_tramite_ran,
        )
        .join(
            models.ProyectoNucleo,
            models.ProyectoNucleo.id_proyecto_nucleo
            == models.TramiteRan.id_proyecto_nucleo,
        )
        .filter(
            models.ProyectoNucleo.id_proyecto == project_id,
            models.ProyectoNucleo.activo.is_(True),
            models.TramiteRan.activo.is_(True),
            models.TramiteRanEvento.activo.is_(True),
        )
        .count()
    )
    return tramites, eventos


def require_catalogs(
    assembly_types: dict[str, Any],
    assembly_contexts: dict[str, Any],
    ran_event_catalog: dict[str, Any],
    result_catalog: dict[str, Any],
) -> None:
    required = [
        ("tipo_asamblea", "anuencia", assembly_types),
        ("tipo_asamblea", "retiro_fondos", assembly_types),
        ("contexto_asamblea", "retiro_fondos", assembly_contexts),
        ("tipo_evento_ran", "ingreso", ran_event_catalog),
        ("tipo_evento_ran", "calificacion", ran_event_catalog),
        ("tipo_evento_ran", "inscripcion", ran_event_catalog),
        ("resultado_convocatoria", "celebrada", result_catalog),
    ]
    for catalog_type, code, catalog in required:
        if code not in catalog:
            raise ImportAbort(f"Falta catálogo activo {catalog_type}/{code}.")


def target_key(item: dict[str, Any]) -> tuple[str, str, tuple[int, ...], str]:
    return (
        item.get("source_key", ""),
        item.get("cop", ""),
        tuple(item.get("source_rows") or [item.get("row")]),
        item.get("target_resolution", ""),
    )


def verify_2br_target(db, item: dict[str, Any]) -> models.Asamblea:
    if item["objective_type"] != "asamblea_anuencia":
        raise ImportAbort(
            f"2F-R sólo admite asamblea_anuencia; recibido {item['objective_type']}."
        )

    target_id = item.get("target_id")
    if target_id is None:
        raise ImportAbort(
            f"Objetivo 2F-R sin id_asamblea para {item.get('source_key')}."
        )

    assembly = db.get(models.Asamblea, target_id)
    if assembly is None or not assembly.activo:
        raise ImportAbort(f"Asamblea objetivo inexistente/inactiva: {target_id}.")

    if assembly.id_proyecto_nucleo != item["pn_id"]:
        raise ImportAbort(
            f"Asamblea {target_id} no corresponde al ProyectoNucleo del trámite."
        )

    obs = assembly.observaciones or ""
    if (
        "Reparación migratoria 2B-R" not in obs
        or item["source_key"] not in obs
    ):
        raise ImportAbort(
            "La Asamblea objetivo no puede acreditarse como creada por 2B-R: "
            f"{item['source_key']} / id_asamblea={target_id}."
        )
    return assembly


def validate_plan_shape(
    db,
    plan: dict[str, Any],
    *,
    allow_create: bool,
) -> tuple[Counter, Counter, Counter, list[dict[str, Any]]]:
    actions, by_objective, events = base.summarize_plan(plan)

    for objective, expected in EXPECTED_SOURCE_ROWS.items():
        actual = plan["source_counts"][objective]
        if actual != expected:
            raise ImportAbort(
                f"Cambió source_rows/{objective}: esperado={expected}, actual={actual}."
            )

    if actions["REVIEW"] != EXPECTED_REVIEW:
        raise ImportAbort(
            f"Cambió REVIEW: esperado={EXPECTED_REVIEW}, actual={actions['REVIEW']}."
        )
    if actions["BLOCKED"] != EXPECTED_BLOCKED:
        raise ImportAbort(
            f"Cambió BLOCKED: esperado={EXPECTED_BLOCKED}, actual={actions['BLOCKED']}."
        )
    if actions["SKIP"] != EXPECTED_SKIP:
        raise ImportAbort(
            f"Cambió SKIP: esperado={EXPECTED_SKIP}, actual={actions['SKIP']}."
        )

    if by_objective[("convenio", "CREATE")] != 0:
        raise ImportAbort("Apareció CREATE de RAN-convenio inesperado.")
    if by_objective[("convenio", "REUSE")] != 73:
        raise ImportAbort(
            "Cambió REUSE/convenio: "
            f"esperado=73, actual={by_objective[('convenio', 'REUSE')]}."
        )
    if by_objective[("convenio", "BLOCKED")] != 14:
        raise ImportAbort(
            "Cambió BLOCKED/convenio: "
            f"esperado=14, actual={by_objective[('convenio', 'BLOCKED')]}."
        )
    if by_objective[("convenio", "REVIEW")] != 0:
        raise ImportAbort("Apareció REVIEW/convenio inesperado.")
    if by_objective[("convenio", "SKIP")] != 1:
        raise ImportAbort("Cambió SKIP/convenio; se esperaba 1.")

    for action, expected in (
        ("CREATE", 0),
        ("REUSE", 3),
        ("REVIEW", 0),
        ("BLOCKED", 0),
        ("SKIP", 0),
    ):
        actual = by_objective[("asamblea_retiro", action)]
        if actual != expected:
            raise ImportAbort(
                f"Cambió {action}/asamblea_retiro: "
                f"esperado={expected}, actual={actual}."
            )

    create_items = [
        item for item in plan["details"] if item["action"] == "CREATE"
    ]
    target_items = [
        item
        for item in plan["details"]
        if item["objective_type"] == "asamblea_anuencia"
        and target_key(item) in EXPECTED_TARGETS
    ]

    observed_target_keys = {target_key(item) for item in target_items}
    if observed_target_keys != EXPECTED_TARGETS:
        raise ImportAbort(
            "Cambió el conjunto exacto de los ocho objetivos 2F-R."
        )
    if len(target_items) != 8:
        raise ImportAbort(
            f"Los objetivos 2F-R deben ser 8; encontrados={len(target_items)}."
        )

    target_actions = Counter(item["action"] for item in target_items)
    if allow_create:
        if target_actions["CREATE"] + target_actions["REUSE"] != 8:
            raise ImportAbort(
                "CREATE+REUSE de objetivos 2F-R debe ser exactamente 8."
            )
        if target_actions["CREATE"] not in {0, 8}:
            raise ImportAbort(
                "Estado parcial 2F-R no permitido: "
                f"CREATE={target_actions['CREATE']} REUSE={target_actions['REUSE']}."
            )
        if target_actions["REUSE"] not in {0, 8}:
            raise ImportAbort(
                "Estado parcial 2F-R no permitido: "
                f"CREATE={target_actions['CREATE']} REUSE={target_actions['REUSE']}."
            )
    else:
        if target_actions["CREATE"] != 0 or target_actions["REUSE"] != 8:
            raise ImportAbort(
                "Postcheck 2F-R exige CREATE=0 / REUSE=8."
            )

    # Todo CREATE del plan tiene que ser uno de los ocho objetivos congelados.
    create_keys = {target_key(item) for item in create_items}
    if not create_keys.issubset(EXPECTED_TARGETS):
        raise ImportAbort(
            "Aparecieron CREATE fuera del plan congelado 2F-R."
        )

    # Acreditar semánticamente las Asambleas de los ocho objetivos.
    for item in target_items:
        if item["action"] in {"CREATE", "REUSE"}:
            verify_2br_target(db, item)

    blocked_anuencia = {
        target_key(item)
        for item in plan["details"]
        if item["objective_type"] == "asamblea_anuencia"
        and item["action"] == "BLOCKED"
    }
    if blocked_anuencia != EXPECTED_REMAINING_ANUENCIA_BLOCKED:
        raise ImportAbort(
            "Cambió el BLOCKED restante de asamblea_anuencia; "
            "se esperaba únicamente LA CUEVA / ADICIONAL / fila 133."
        )

    if allow_create:
        expected_anuencia_reuse = 47 + target_actions["REUSE"]
        expected_anuencia_create = target_actions["CREATE"]
    else:
        expected_anuencia_reuse = 55
        expected_anuencia_create = 0

    checks = {
        "CREATE": expected_anuencia_create,
        "REUSE": expected_anuencia_reuse,
        "REVIEW": 6,
        "BLOCKED": 1,
        "SKIP": 1,
    }
    for action, expected in checks.items():
        actual = by_objective[("asamblea_anuencia", action)]
        if actual != expected:
            raise ImportAbort(
                f"Cambió {action}/asamblea_anuencia: "
                f"esperado={expected}, actual={actual}."
            )

    # Contar sólo los eventos de los ocho objetivos de reparación.
    repair_events = Counter()
    for item in target_items:
        candidate = item.get("candidate")
        if item["action"] in {"CREATE", "REUSE"} and candidate:
            for event in candidate["events"]:
                repair_events[event["type_code"]] += 1

    for event_type, expected in EXPECTED_NEW_EVENTS.items():
        if repair_events[event_type] != expected:
            raise ImportAbort(
                f"Cambió eventos 2F-R/{event_type}: "
                f"esperado={expected}, actual={repair_events[event_type]}."
            )

    return actions, by_objective, events, target_items


def create_tramite(
    db,
    *,
    item: dict[str, Any],
    ran_event_catalog: dict[str, Any],
    actor_id: int,
    excel_name: str,
) -> models.TramiteRan:
    candidate = item["candidate"]

    matches = base.existing_exact_matches(
        db=db,
        objective_type=item["objective_type"],
        objective_id=item["target_id"],
        pn_id=item["pn_id"],
        candidate=candidate,
        event_catalog=ran_event_catalog,
    )
    if matches:
        raise ImportAbort(
            "El objetivo marcado CREATE ya tiene un TramiteRan de firma exacta: "
            f"{item['source_key']} / filas={item['source_rows']}."
        )

    tramite = models.TramiteRan(
        id_proyecto_nucleo=item["pn_id"],
        id_nucleo=None,
        id_asamblea=item["target_id"],
        id_convenio=None,
        id_orv=None,
        fecha_programada_ingreso=candidate["fecha_programada_ingreso"],
        referencia_expediente=None,
        creado_por=actor_id,
        observaciones=(
            f"Reparación migratoria 2F-R desde {excel_name}; "
            f"hoja {SHEET}; "
            f"filas {','.join(map(str, item['source_rows']))}; "
            f"objetivo=asamblea_anuencia; "
            f"resolucion={item['target_resolution']}; "
            f"COP={item.get('cop') or '-'}."
        ),
    )
    db.add(tramite)
    db.flush()

    for event in candidate["events"]:
        type_option = ran_event_catalog[event["type_code"]]
        row = models.TramiteRanEvento(
            id_tramite_ran=tramite.id_tramite_ran,
            ordinal=event["ordinal"],
            id_tipo_evento=type_option.id_catalogo_opcion,
            fecha_evento=event["fecha_evento"],
            numero_solicitud=event["numero_solicitud"],
            resultado=event["resultado"],
            calificacion=event["calificacion"],
            folio_referencia=event["folio_referencia"],
            id_documento=None,
            creado_por=actor_id,
            observaciones=(
                f"Reparación migratoria 2F-R desde {excel_name}; "
                f"hoja {SHEET}; "
                f"filas {','.join(map(str, item['source_rows']))}; "
                f"evento={event['type_code']}."
            ),
        )
        db.add(row)

    db.flush()
    return tramite


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)

    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--confirmar", action="store_true")

    parser.add_argument("--actor-id", type=int, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    args = parser.parse_args()

    if DB_NAME != EXPECTED_DB:
        raise ImportAbort(
            f"Protección: DB_NAME debe ser {EXPECTED_DB!r}; recibido {DB_NAME!r}."
        )

    if not args.excel.exists():
        raise ImportAbort(f"No existe el Excel: {args.excel}")

    digest = base.sha256_file(args.excel)
    if digest != EXPECTED_SHA256:
        raise ImportAbort(
            "El SHA-256 no coincide con el Excel aprobado. "
            f"Esperado {EXPECTED_SHA256}; recibido {digest}."
        )

    workbook = load_workbook(args.excel, data_only=True, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise ImportAbort(f"Falta la hoja {SHEET!r}.")
    sheet = workbook[SHEET]

    args.output_dir.mkdir(parents=True, exist_ok=True)
    mode_name = "confirmar" if args.confirmar else "dry-run"
    report_path = (
        args.output_dir
        / f"importacion_reparacion_ran_2f_r_{mode_name}.json"
    )

    db = SessionLocal()

    try:
        current_db = db.execute(text("SELECT current_database()")).scalar()
        current_schema = base.schema_version(db)

        if current_db != EXPECTED_DB or current_schema != EXPECTED_SCHEMA:
            raise ImportAbort(
                f"Destino inesperado: BD={current_db!r}, schema={current_schema!r}."
            )

        actor = (
            db.query(models.Usuario)
            .filter(
                models.Usuario.id_usuario == args.actor_id,
                models.Usuario.activo.is_(True),
            )
            .one_or_none()
        )
        if actor is None:
            raise ImportAbort(
                f"Actor {args.actor_id} no existe o no está activo."
            )

        project = (
            db.query(models.Proyecto)
            .filter(
                models.Proyecto.clave_proyecto == PROJECT_KEY,
                models.Proyecto.activo.is_(True),
            )
            .one_or_none()
        )
        if project is None:
            raise ImportAbort(f"No existe proyecto activo {PROJECT_KEY!r}.")

        before_tramites, before_eventos = project_ran_counts(
            db, project.id_proyecto
        )
        if (before_tramites, before_eventos) not in {
            (123, 208),
            (131, 223),
        }:
            raise ImportAbort(
                "Estado persistido 2F-R inesperado: "
                f"TramiteRan={before_tramites}, Eventos={before_eventos}; "
                "se esperaba 123/208 o 131/223."
            )

        pn_index = base.build_project_nucleus_index(db, project)
        cop_catalog = base.catalog_by_code(db, "tipo_cop_operativo")
        assembly_types = base.catalog_by_code(db, "tipo_asamblea")
        assembly_contexts = base.catalog_by_code(db, "contexto_asamblea")
        ran_event_catalog = base.catalog_by_code(db, "tipo_evento_ran")
        result_catalog = base.catalog_by_code(db, "resultado_convocatoria")

        require_catalogs(
            assembly_types,
            assembly_contexts,
            ran_event_catalog,
            result_catalog,
        )

        plan = base.build_plan(
            db=db,
            sheet=sheet,
            pn_index=pn_index,
            cop_catalog=cop_catalog,
            assembly_types=assembly_types,
            assembly_contexts=assembly_contexts,
            ran_event_catalog=ran_event_catalog,
            result_catalog=result_catalog,
        )

        (
            pre_actions,
            pre_objectives,
            pre_events,
            target_items,
        ) = validate_plan_shape(db, plan, allow_create=True)

        target_action_counts = Counter(
            item["action"] for item in target_items
        )

        expected_state = (
            (123, 208)
            if target_action_counts["CREATE"] == 8
            else (131, 223)
        )
        if (before_tramites, before_eventos) != expected_state:
            raise ImportAbort(
                "Los conteos persistidos no concuerdan con CREATE/REUSE 2F-R: "
                f"BD={before_tramites}/{before_eventos}, "
                f"objetivos CREATE={target_action_counts['CREATE']} "
                f"REUSE={target_action_counts['REUSE']}."
            )

        db.execute(
            text("SELECT set_config('app.current_user_id', :id, true)"),
            {"id": str(args.actor_id)},
        )

        runtime = Counter()
        imported = []

        for item in sorted(
            target_items,
            key=lambda x: (
                min(x.get("source_rows") or [x.get("row")]),
                x["source_key"],
            ),
        ):
            candidate = item["candidate"]

            matches = base.existing_exact_matches(
                db=db,
                objective_type=item["objective_type"],
                objective_id=item["target_id"],
                pn_id=item["pn_id"],
                candidate=candidate,
                event_catalog=ran_event_catalog,
            )

            if len(matches) > 1:
                raise ImportAbort(
                    "Hay múltiples TramiteRan con la firma exacta del objetivo "
                    f"2F-R {item['source_key']} / filas={item['source_rows']}."
                )

            if len(matches) == 1:
                tramite = matches[0]
                action = "REUSE"
                runtime["REUSE"] += 1
            else:
                verify_2br_target(db, item)
                tramite = create_tramite(
                    db,
                    item=item,
                    ran_event_catalog=ran_event_catalog,
                    actor_id=args.actor_id,
                    excel_name=args.excel.name,
                )
                action = "CREATE"
                runtime["CREATE"] += 1

            imported.append({
                "action": action,
                "id_tramite_ran": tramite.id_tramite_ran,
                **base.serialize_item(item),
            })

        if runtime["CREATE"] + runtime["REUSE"] != 8:
            raise ImportAbort(
                "Runtime 2F-R no procesó exactamente ocho objetivos."
            )
        if runtime["CREATE"] not in {0, 8} or runtime["REUSE"] not in {0, 8}:
            raise ImportAbort(
                "Runtime 2F-R quedó en estado parcial: "
                f"CREATE={runtime['CREATE']} REUSE={runtime['REUSE']}."
            )

        # Forzar constraints antes de commit/rollback.
        db.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))

        in_tx_tramites, in_tx_eventos = project_ran_counts(
            db, project.id_proyecto
        )
        if (
            in_tx_tramites != EXPECTED_FINAL_TRAMITES
            or in_tx_eventos != EXPECTED_FINAL_EVENTOS
        ):
            raise ImportAbort(
                "Postcheck intratransacción de totales RAN falló: "
                f"TramiteRan={in_tx_tramites} / "
                f"Eventos={in_tx_eventos}; "
                f"esperado={EXPECTED_FINAL_TRAMITES}/{EXPECTED_FINAL_EVENTOS}."
            )

        # Rehacer completamente el plan dentro de la misma transacción.
        plan_after = base.build_plan(
            db=db,
            sheet=sheet,
            pn_index=pn_index,
            cop_catalog=cop_catalog,
            assembly_types=assembly_types,
            assembly_contexts=assembly_contexts,
            ran_event_catalog=ran_event_catalog,
            result_catalog=result_catalog,
        )

        (
            post_actions,
            post_objectives,
            post_events,
            post_targets,
        ) = validate_plan_shape(db, plan_after, allow_create=False)

        for event_type, expected in EXPECTED_FINAL_EVENTS.items():
            actual = post_events[event_type]
            if actual != expected:
                raise ImportAbort(
                    f"Eventos finales/{event_type}: "
                    f"esperado={expected}, actual={actual}."
                )

        if post_actions["CREATE"] != 0 or post_actions["REUSE"] != 131:
            raise ImportAbort(
                "Idempotencia intratransacción falló: "
                f"CREATE={post_actions['CREATE']} REUSE={post_actions['REUSE']}."
            )

        verified = []
        for item in sorted(
            post_targets,
            key=lambda x: (
                min(x.get("source_rows") or [x.get("row")]),
                x["source_key"],
            ),
        ):
            if item["action"] != "REUSE":
                raise ImportAbort("Postcheck: objetivo 2F-R no quedó en REUSE.")

            matches = base.existing_exact_matches(
                db=db,
                objective_type=item["objective_type"],
                objective_id=item["target_id"],
                pn_id=item["pn_id"],
                candidate=item["candidate"],
                event_catalog=ran_event_catalog,
            )
            if len(matches) != 1:
                raise ImportAbort(
                    "Postcheck: firma 2F-R no resuelve exactamente un TramiteRan."
                )

            verified.append({
                "source_key": item["source_key"],
                "cop": item["cop"],
                "source_rows": item["source_rows"],
                "id_asamblea": item["target_id"],
                "id_tramite_ran": matches[0].id_tramite_ran,
                "target_resolution": item["target_resolution"],
                "events": [
                    event["type_code"] for event in item["candidate"]["events"]
                ],
            })

        report = {
            "mode": mode_name,
            "database": current_db,
            "schema": current_schema,
            "actor": {
                "id_usuario": actor.id_usuario,
                "correo": actor.correo,
            },
            "project": PROJECT_KEY,
            "excel": args.excel.name,
            "sha256": digest,
            "before_persisted": {
                "tramite_ran": before_tramites,
                "tramite_ran_evento": before_eventos,
            },
            "targets_before": {
                "CREATE": target_action_counts["CREATE"],
                "REUSE": target_action_counts["REUSE"],
            },
            "runtime": {
                "CREATE": runtime["CREATE"],
                "REUSE": runtime["REUSE"],
            },
            "after_in_transaction": {
                "tramite_ran": in_tx_tramites,
                "tramite_ran_evento": in_tx_eventos,
                "CREATE": post_actions["CREATE"],
                "REUSE": post_actions["REUSE"],
                "REVIEW": post_actions["REVIEW"],
                "BLOCKED": post_actions["BLOCKED"],
                "SKIP": post_actions["SKIP"],
            },
            "verified_targets": verified,
            "remaining_blocked": [
                base.serialize_item(item)
                for item in plan_after["details"]
                if item["action"] == "BLOCKED"
            ],
            "reviews": [
                base.serialize_item(item)
                for item in plan_after["details"]
                if item["action"] == "REVIEW"
            ],
        }

        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

        print("=== IMPORTACIÓN REPARACIÓN 2F-R — TRÁMITES RAN ===")
        print(f"Modo: {mode_name}")
        print(f"BD: {current_db} / schema {current_schema}")
        print(f"Actor: {actor.id_usuario} / {actor.correo}")
        print(f"Proyecto: {PROJECT_KEY}")
        print(f"SHA-256: {digest}")
        print("Persistencia al inicio:")
        print(f"  TramiteRan:        {before_tramites}")
        print(f"  TramiteRanEvento:  {before_eventos}")
        print("Objetivos 2F-R al inicio:")
        print(f"  CREATE   {target_action_counts['CREATE']}")
        print(f"  REUSE    {target_action_counts['REUSE']}")
        print("Runtime 2F-R:")
        print(f"  CREATE   {runtime['CREATE']}")
        print(f"  REUSE    {runtime['REUSE']}")
        print("Validación intratransacción global:")
        print(f"  TramiteRan:        {in_tx_tramites}")
        print(f"  TramiteRanEvento:  {in_tx_eventos}")
        print(f"  CREATE             {post_actions['CREATE']}")
        print(f"  REUSE              {post_actions['REUSE']}")
        print(f"  REVIEW             {post_actions['REVIEW']}")
        print(f"  BLOCKED            {post_actions['BLOCKED']}")
        print(f"  SKIP               {post_actions['SKIP']}")
        print("Objetivos 2F-R verificados:")
        for item in verified:
            print(
                f"  {item['source_key']} | {item['cop']} | "
                f"filas={','.join(map(str, item['source_rows']))} | "
                f"id_asamblea={item['id_asamblea']} -> "
                f"id_tramite_ran={item['id_tramite_ran']} | "
                f"{item['target_resolution']} | "
                f"eventos={','.join(item['events'])}"
            )
        print("Permanecen deliberadamente:")
        print("  REVIEW   6")
        print("  BLOCKED  15 (1 asamblea_anuencia + 14 convenio)")
        print("  SKIP     2")

        if args.confirmar:
            db.commit()
            print("COMMIT realizado")
        else:
            db.rollback()
            print("ROLLBACK realizado (dry-run)")

        print(f"Reporte: {report_path}")
        return 0

    except IntegrityError as exc:
        db.rollback()
        print("=== IMPORTACIÓN 2F-R ABORTADA ===", file=sys.stderr)
        print(f"IntegrityError: {exc}", file=sys.stderr)
        return 2
    except (ImportAbort, base.ImportAbort) as exc:
        db.rollback()
        print("=== IMPORTACIÓN 2F-R ABORTADA ===", file=sys.stderr)
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
