#!/usr/bin/env python3
"""Preflight read-only 2F-R: reevaluar trámites RAN tras 2B-R / 2D-R.

Objetivo
========
Recalcular la capa 2F con la lógica original, sobre el estado actual de
db_carga_excel, sin escribir nada. El propósito es separar con evidencia:

- trámites originales ya importados que siguen resolviendo exactamente (REUSE);
- casos que estaban BLOCKED y ahora se volvieron deterministas (CREATE);
- casos que siguen BLOCKED;
- casos que ahora requieren REVIEW;
- SKIP que deben permanecer como tales.

Este script NO importa ni modifica TramiteRan / TramiteRanEvento.

Base histórica congelada de 2F
==============================
- 88 filas fuente RAN de asamblea de anuencia
- 88 filas fuente RAN de convenio
- 3 filas fuente RAN de asamblea de retiro
- 123 trámites deterministas ya importados
- 208 eventos ya importados:
    123 ingreso
    9 calificación
    76 inscripción
- 6 REVIEW
- 23 BLOCKED
    9 asamblea_anuencia
    14 convenio
- 2 SKIP
- por objetivo determinista:
    47 asamblea_anuencia
    73 convenio
    3 asamblea_retiro

Cambios posteriores permitidos
===============================
2B-R creó Asambleas de anuencia únicamente en seis grupos auditados:
- MEXICO|COYOTEPEC|COYOTEPEC
- HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE
- HIDALGO|CHAPANTONGO|JUCHITLAN
- HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ
- MEXICO|JILOTEPEC|SANTIAGO OXTHOC
- QUERETARO|SAN JUAN DEL RIO|PASO DE MATA

2D-R NO creó Convenios; sólo completó id_asamblea_autorizacion en cuatro
Convenios ya existentes. Por tanto, un BLOCKED de RAN-convenio no debe
volverse CREATE como consecuencia de 2D-R.

Guardas
=======
- sólo DB_NAME=db_carga_excel
- schema 018
- SHA-256 exacto del Excel aprobado
- proyecto MEX-QRO
- fuente RAN conserva 88 / 88 / 3 filas
- existen exactamente los 123 TramiteRan y 208 TramiteRanEvento históricos
- los 123 trámites históricos deben seguir resolviendo como REUSE
- los 73 trámites de convenio deben seguir REUSE
- los 14 BLOCKED de convenio deben seguir BLOCKED
- los 3 trámites de retiro deben seguir REUSE
- cualquier CREATE nuevo sólo puede ser asamblea_anuencia y sólo en uno de
  los seis grupos reparados por 2B-R
- una Asamblea objetivo de un CREATE debe acreditarse como creada por 2B-R
- no se cambian conteos esperados para hacer pasar el preflight

La salida de este preflight se usa para congelar el plan 2F-R antes de
construir cualquier importador.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app import models  # noqa: E402
from app.database import DB_NAME, SessionLocal  # noqa: E402
import importar_ran_colectivos as base  # noqa: E402


EXPECTED_DB = "db_carga_excel"
EXPECTED_SCHEMA = "018"
EXPECTED_SHA256 = "bac01c25e9ff963f524fa8a48d69e74d92a0d36578731b1eca91f0e5a3573c3f"
PROJECT_KEY = "MEX-QRO"
SHEET = "INFORME M-Q"

EXPECTED_SOURCE_ROWS = {
    "asamblea_anuencia": 88,
    "convenio": 88,
    "asamblea_retiro": 3,
}

# Estado ya persistido por la capa 2F original.
EXPECTED_EXISTING_TRAMITES = 123
EXPECTED_EXISTING_EVENTOS = 208

# Los 123 deterministas originales deben seguir resolviendo exactamente.
EXPECTED_REUSE_ORIGINAL = {
    "asamblea_anuencia": 47,
    "convenio": 73,
    "asamblea_retiro": 3,
}

# 2D-R no crea Convenios, así que estos 14 no deben cambiar por esa reparación.
EXPECTED_CONVENIO_BLOCKED = 14
EXPECTED_CONVENIO_REVIEW = 0
EXPECTED_CONVENIO_SKIP = 1

# Asamblea de retiro no fue modificada por 2B-R/2D-R.
EXPECTED_RETIRO_REVIEW = 0
EXPECTED_RETIRO_BLOCKED = 0
EXPECTED_RETIRO_SKIP = 0

# SKIP histórico de asamblea de anuencia (incluye el PN no importable).
EXPECTED_ANUENCIA_SKIP = 1

# Los 9 BLOCKED originales de asamblea de anuencia deben repartirse ahora entre
# CREATE / BLOCKED / REVIEW adicional. Los 6 REVIEW históricos deben preservarse.
EXPECTED_ANUENCIA_PREVIOUS_BLOCKED = 9
EXPECTED_ANUENCIA_BASE_REVIEW = 6

REPAIRED_2BR_KEYS = {
    "MEXICO|COYOTEPEC|COYOTEPEC",
    "HIDALGO|TULA DE ALLENDE|TULA DE ALLENDE",
    "HIDALGO|CHAPANTONGO|JUCHITLAN",
    "HIDALGO|CHAPANTONGO|SAN JOSE EL MARQUEZ",
    "MEXICO|JILOTEPEC|SANTIAGO OXTHOC",
    "QUERETARO|SAN JUAN DEL RIO|PASO DE MATA",
}


class PreflightAbort(RuntimeError):
    pass


def active_project_ran_counts(db, project_id: int) -> tuple[int, int]:
    tramite_ids = (
        db.query(models.TramiteRan.id_tramite_ran)
        .join(
            models.ProyectoNucleo,
            models.ProyectoNucleo.id_proyecto_nucleo
            == models.TramiteRan.id_proyecto_nucleo,
        )
        .filter(
            models.ProyectoNucleo.id_proyecto == project_id,
            models.ProyectoNucleo.activo.is_(True),
            models.TramiteRan.activo.is_(True),
        )
        .subquery()
    )

    tramites = db.query(models.TramiteRan).filter(
        models.TramiteRan.id_tramite_ran.in_(db.query(tramite_ids.c.id_tramite_ran))
    ).count()

    eventos = (
        db.query(models.TramiteRanEvento)
        .filter(
            models.TramiteRanEvento.id_tramite_ran.in_(
                db.query(tramite_ids.c.id_tramite_ran)
            ),
            models.TramiteRanEvento.activo.is_(True),
        )
        .count()
    )

    return tramites, eventos


def require_catalogs(
    assembly_types: dict[str, Any],
    assembly_contexts: dict[str, Any],
    ran_event_catalog: dict[str, Any],
    result_catalog: dict[str, Any],
) -> None:
    required = [
        ("tipo_asamblea", "anuencia", assembly_types),
        ("tipo_asamblea", "retiro_fondos", assembly_types),
        ("contexto_asamblea", "retiro_fondos", assembly_contexts),
        ("tipo_evento_ran", "ingreso", ran_event_catalog),
        ("tipo_evento_ran", "calificacion", ran_event_catalog),
        ("tipo_evento_ran", "inscripcion", ran_event_catalog),
        ("resultado_convocatoria", "celebrada", result_catalog),
    ]
    for catalog_type, code, catalog in required:
        if code not in catalog:
            raise PreflightAbort(
                f"Falta catálogo activo {catalog_type}/{code}."
            )


def event_counts_for_items(items: list[dict[str, Any]]) -> Counter:
    counts = Counter()
    for item in items:
        candidate = item.get("candidate")
        if not candidate:
            continue
        for event in candidate["events"]:
            counts[event["type_code"]] += 1
    return counts


def serializable_detail(item: dict[str, Any]) -> dict[str, Any]:
    candidate = item.get("candidate")
    events = []

    if candidate:
        for event in candidate["events"]:
            events.append({
                "ordinal": event["ordinal"],
                "tipo": event["type_code"],
                "fecha_evento": (
                    event["fecha_evento"].isoformat()
                    if event["fecha_evento"] is not None
                    else None
                ),
                "numero_solicitud": event["numero_solicitud"],
                "resultado": event["resultado"],
                "calificacion": event["calificacion"],
                "folio_referencia": event["folio_referencia"],
            })

    return {
        "action": item["action"],
        "objective_type": item["objective_type"],
        "source_rows": item.get("source_rows") or [item.get("row")],
        "source_key": item.get("source_key", ""),
        "cop": item.get("cop", ""),
        "target_id": item.get("target_id"),
        "target_resolution": item.get("target_resolution", ""),
        "dedup_count": item.get("dedup_count", 1),
        "scheduled_raw": item.get("scheduled_raw", ""),
        "ingreso_raw": item.get("ingreso_raw", ""),
        "request_raw": item.get("request_raw", ""),
        "calificacion_raw": item.get("calificacion_raw", ""),
        "inscripcion_raw": item.get("inscripcion_raw", ""),
        "fecha_programada_ingreso": (
            candidate["fecha_programada_ingreso"].isoformat()
            if candidate and candidate["fecha_programada_ingreso"]
            else None
        ),
        "eventos": events,
        "detail": item.get("detail", ""),
    }


def validate_current_plan(db, plan: dict[str, Any]) -> dict[str, Any]:
    source_counts = plan["source_counts"]
    actions, by_objective, _events = base.summarize_plan(plan)

    for objective, expected in EXPECTED_SOURCE_ROWS.items():
        actual = source_counts[objective]
        if actual != expected:
            raise PreflightAbort(
                f"Cambió source_rows/{objective}: "
                f"esperado={expected}, actual={actual}."
            )

    # Los 123 deterministas originales deben seguir resolviendo como REUSE.
    original_reuse_total = sum(EXPECTED_REUSE_ORIGINAL.values())
    if actions["REUSE"] != original_reuse_total:
        raise PreflightAbort(
            "Los trámites 2F históricos ya no resuelven exactamente como antes: "
            f"REUSE esperado={original_reuse_total}, actual={actions['REUSE']}."
        )

    for objective, expected in EXPECTED_REUSE_ORIGINAL.items():
        actual = by_objective[(objective, "REUSE")]
        if actual != expected:
            raise PreflightAbort(
                f"Cambió REUSE/{objective}: esperado={expected}, actual={actual}."
            )

    # Convenio: 2D-R no creó padres nuevos. Si algo cambia aquí, no es una
    # consecuencia legítima de la reparación de Asamblea.
    if by_objective[("convenio", "CREATE")] != 0:
        raise PreflightAbort(
            "Aparecieron CREATE de RAN-convenio, pero 2D-R no creó Convenios."
        )
    if by_objective[("convenio", "BLOCKED")] != EXPECTED_CONVENIO_BLOCKED:
        raise PreflightAbort(
            "Cambió BLOCKED/convenio: "
            f"esperado={EXPECTED_CONVENIO_BLOCKED}, "
            f"actual={by_objective[('convenio', 'BLOCKED')]}."
        )
    if by_objective[("convenio", "REVIEW")] != EXPECTED_CONVENIO_REVIEW:
        raise PreflightAbort(
            "Cambió REVIEW/convenio: "
            f"esperado={EXPECTED_CONVENIO_REVIEW}, "
            f"actual={by_objective[('convenio', 'REVIEW')]}."
        )
    if by_objective[("convenio", "SKIP")] != EXPECTED_CONVENIO_SKIP:
        raise PreflightAbort(
            "Cambió SKIP/convenio: "
            f"esperado={EXPECTED_CONVENIO_SKIP}, "
            f"actual={by_objective[('convenio', 'SKIP')]}."
        )

    # Retiro tampoco fue tocado.
    for action, expected in (
        ("CREATE", 0),
        ("REVIEW", EXPECTED_RETIRO_REVIEW),
        ("BLOCKED", EXPECTED_RETIRO_BLOCKED),
        ("SKIP", EXPECTED_RETIRO_SKIP),
    ):
        actual = by_objective[("asamblea_retiro", action)]
        if actual != expected:
            raise PreflightAbort(
                f"Cambió {action}/asamblea_retiro: "
                f"esperado={expected}, actual={actual}."
            )

    if by_objective[("asamblea_anuencia", "SKIP")] != EXPECTED_ANUENCIA_SKIP:
        raise PreflightAbort(
            "Cambió SKIP/asamblea_anuencia: "
            f"esperado={EXPECTED_ANUENCIA_SKIP}, "
            f"actual={by_objective[('asamblea_anuencia', 'SKIP')]}."
        )

    # Los 6 REVIEW históricos de anuencia deben seguir presentes. Los 9 BLOCKED
    # históricos pueden repartirse entre CREATE, BLOCKED o REVIEW adicional.
    anuencia_create = by_objective[("asamblea_anuencia", "CREATE")]
    anuencia_blocked = by_objective[("asamblea_anuencia", "BLOCKED")]
    anuencia_review = by_objective[("asamblea_anuencia", "REVIEW")]

    if anuencia_review < EXPECTED_ANUENCIA_BASE_REVIEW:
        raise PreflightAbort(
            "Desaparecieron REVIEW históricos de asamblea_anuencia: "
            f"mínimo esperado={EXPECTED_ANUENCIA_BASE_REVIEW}, "
            f"actual={anuencia_review}."
        )

    redistributed = (
        anuencia_create
        + anuencia_blocked
        + (anuencia_review - EXPECTED_ANUENCIA_BASE_REVIEW)
    )
    if redistributed != EXPECTED_ANUENCIA_PREVIOUS_BLOCKED:
        raise PreflightAbort(
            "Los 9 BLOCKED históricos de asamblea_anuencia no se reconciliaron "
            "de forma conservadora: "
            f"CREATE={anuencia_create}, BLOCKED={anuencia_blocked}, "
            f"REVIEW_adicional={anuencia_review - EXPECTED_ANUENCIA_BASE_REVIEW}, "
            f"total={redistributed}."
        )

    if actions["SKIP"] != 2:
        raise PreflightAbort(
            f"Cambió SKIP total de 2F: esperado=2, actual={actions['SKIP']}."
        )

    create_items = [
        item for item in plan["details"] if item["action"] == "CREATE"
    ]

    for item in create_items:
        if item["objective_type"] != "asamblea_anuencia":
            raise PreflightAbort(
                "CREATE nuevo fuera de asamblea_anuencia: "
                f"{item['objective_type']} / {item.get('source_key')}."
            )

        source_key = item.get("source_key", "")
        if source_key not in REPAIRED_2BR_KEYS:
            raise PreflightAbort(
                "CREATE de RAN apareció fuera de los seis grupos reparados por 2B-R: "
                f"{source_key}."
            )

        target_id = item.get("target_id")
        if target_id is None:
            raise PreflightAbort(
                f"CREATE sin id_asamblea objetivo para {source_key}."
            )

        assembly = db.get(models.Asamblea, target_id)
        if assembly is None or not assembly.activo:
            raise PreflightAbort(
                f"Asamblea objetivo inexistente/inactiva: id={target_id}."
            )

        obs = assembly.observaciones or ""
        if (
            "Reparación migratoria 2B-R" not in obs
            or source_key not in obs
        ):
            raise PreflightAbort(
                "El CREATE nuevo no apunta a una Asamblea acreditable como 2B-R: "
                f"{source_key} / id_asamblea={target_id}."
            )

    new_event_counts = event_counts_for_items(create_items)

    return {
        "actions": actions,
        "by_objective": by_objective,
        "create_items": create_items,
        "new_event_counts": new_event_counts,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    args = parser.parse_args()

    if DB_NAME != EXPECTED_DB:
        raise PreflightAbort(
            f"Protección: DB_NAME debe ser {EXPECTED_DB!r}; recibido {DB_NAME!r}."
        )

    if not args.excel.exists():
        raise PreflightAbort(f"No existe el Excel: {args.excel}")

    digest = base.sha256_file(args.excel)
    if digest != EXPECTED_SHA256:
        raise PreflightAbort(
            "El SHA-256 no coincide con el Excel aprobado. "
            f"Esperado {EXPECTED_SHA256}; recibido {digest}."
        )

    workbook = load_workbook(args.excel, data_only=True, read_only=False)
    if SHEET not in workbook.sheetnames:
        raise PreflightAbort(f"Falta la hoja {SHEET!r}.")
    sheet = workbook[SHEET]

    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.output_dir / "preflight_reparacion_ran_2f_r.csv"
    json_path = args.output_dir / "preflight_reparacion_ran_2f_r.json"

    db = SessionLocal()

    try:
        current_db = db.execute(text("SELECT current_database()")).scalar()
        current_schema = base.schema_version(db)

        if current_db != EXPECTED_DB or current_schema != EXPECTED_SCHEMA:
            raise PreflightAbort(
                f"Destino inesperado: BD={current_db!r}, schema={current_schema!r}."
            )

        project = (
            db.query(models.Proyecto)
            .filter(
                models.Proyecto.clave_proyecto == PROJECT_KEY,
                models.Proyecto.activo.is_(True),
            )
            .one_or_none()
        )
        if project is None:
            raise PreflightAbort(f"No existe proyecto activo {PROJECT_KEY!r}.")

        tramites, eventos = active_project_ran_counts(db, project.id_proyecto)
        if tramites != EXPECTED_EXISTING_TRAMITES:
            raise PreflightAbort(
                "Cambió el total persistido de TramiteRan previo a 2F-R: "
                f"esperado={EXPECTED_EXISTING_TRAMITES}, actual={tramites}."
            )
        if eventos != EXPECTED_EXISTING_EVENTOS:
            raise PreflightAbort(
                "Cambió el total persistido de TramiteRanEvento previo a 2F-R: "
                f"esperado={EXPECTED_EXISTING_EVENTOS}, actual={eventos}."
            )

        pn_index = base.build_project_nucleus_index(db, project)
        cop_catalog = base.catalog_by_code(db, "tipo_cop_operativo")
        assembly_types = base.catalog_by_code(db, "tipo_asamblea")
        assembly_contexts = base.catalog_by_code(db, "contexto_asamblea")
        ran_event_catalog = base.catalog_by_code(db, "tipo_evento_ran")
        result_catalog = base.catalog_by_code(db, "resultado_convocatoria")

        require_catalogs(
            assembly_types,
            assembly_contexts,
            ran_event_catalog,
            result_catalog,
        )

        plan = base.build_plan(
            db,
            sheet,
            pn_index,
            cop_catalog,
            assembly_types,
            assembly_contexts,
            ran_event_catalog,
            result_catalog,
        )

        validated = validate_current_plan(db, plan)
        actions = validated["actions"]
        by_objective = validated["by_objective"]
        create_items = validated["create_items"]
        new_events = validated["new_event_counts"]

        details = [
            serializable_detail(item)
            for item in sorted(
                plan["details"],
                key=lambda x: (
                    x["objective_type"],
                    (x.get("source_rows") or [x.get("row", 0)])[0],
                    x["action"],
                ),
            )
        ]

        fields = [
            "action",
            "objective_type",
            "source_rows",
            "source_key",
            "cop",
            "target_id",
            "target_resolution",
            "dedup_count",
            "scheduled_raw",
            "ingreso_raw",
            "request_raw",
            "calificacion_raw",
            "inscripcion_raw",
            "fecha_programada_ingreso",
            "eventos",
            "detail",
        ]

        with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            for item in details:
                writer.writerow({
                    **item,
                    "source_rows": ";".join(map(str, item["source_rows"])),
                    "eventos": " | ".join(
                        (
                            f"{e['ordinal']}:{e['tipo']}"
                            f":fecha={e['fecha_evento'] or '-'}"
                            f":sol={e['numero_solicitud'] or '-'}"
                            f":resultado={e['resultado'] or '-'}"
                            f":calif={e['calificacion'] or '-'}"
                        )
                        for e in item["eventos"]
                    ),
                })

        report = {
            "database": current_db,
            "schema": current_schema,
            "project": PROJECT_KEY,
            "excel": args.excel.name,
            "sha256": digest,
            "persisted_before_repair": {
                "tramite_ran": tramites,
                "tramite_ran_evento": eventos,
            },
            "source_rows": {
                key: plan["source_counts"][key]
                for key in EXPECTED_SOURCE_ROWS
            },
            "actions": {
                action: actions[action]
                for action in ("CREATE", "REUSE", "REVIEW", "BLOCKED", "SKIP")
            },
            "by_objective": {
                objective: {
                    action: by_objective[(objective, action)]
                    for action in ("CREATE", "REUSE", "REVIEW", "BLOCKED", "SKIP")
                }
                for objective in (
                    "asamblea_anuencia",
                    "convenio",
                    "asamblea_retiro",
                )
            },
            "new_create_events": {
                event: new_events[event]
                for event in ("ingreso", "calificacion", "inscripcion")
            },
            "new_create_details": [
                serializable_detail(item) for item in create_items
            ],
            "remaining_blocked": [
                serializable_detail(item)
                for item in plan["details"]
                if item["action"] == "BLOCKED"
            ],
            "reviews": [
                serializable_detail(item)
                for item in plan["details"]
                if item["action"] == "REVIEW"
            ],
        }

        json_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

        print("=== PREFLIGHT REPARACIÓN 2F-R — TRÁMITES RAN ===")
        print(f"BD: {current_db} / schema {current_schema}")
        print(f"Proyecto: {PROJECT_KEY}")
        print(f"SHA-256: {digest}")
        print("Persistencia 2F previa:")
        print(f"  TramiteRan:        {tramites}")
        print(f"  TramiteRanEvento:  {eventos}")
        print("Filas fuente preservadas:")
        for objective in (
            "asamblea_anuencia",
            "convenio",
            "asamblea_retiro",
        ):
            print(
                f"  {objective:<20} "
                f"{plan['source_counts'][objective]}"
            )

        print("Acciones recalculadas después de 2B-R / 2D-R:")
        for action in ("CREATE", "REUSE", "REVIEW", "BLOCKED", "SKIP"):
            print(f"  {action:<8} {actions[action]}")

        print("Por objetivo:")
        for objective in (
            "asamblea_anuencia",
            "convenio",
            "asamblea_retiro",
        ):
            print(
                f"  {objective}: "
                f"CREATE={by_objective[(objective, 'CREATE')]} "
                f"REUSE={by_objective[(objective, 'REUSE')]} "
                f"REVIEW={by_objective[(objective, 'REVIEW')]} "
                f"BLOCKED={by_objective[(objective, 'BLOCKED')]} "
                f"SKIP={by_objective[(objective, 'SKIP')]}"
            )

        print("Nuevos candidatos deterministas desbloqueados:")
        if not create_items:
            print("  Ninguno")
        else:
            for item in create_items:
                candidate = item["candidate"]
                event_types = ",".join(
                    event["type_code"] for event in candidate["events"]
                )
                print(
                    f"  {item['source_key']} | {item['cop']} | "
                    f"filas={','.join(map(str, item.get('source_rows') or [item['row']]))} | "
                    f"id_asamblea={item['target_id']} | "
                    f"{item['target_resolution']} | eventos={event_types or '-'}"
                )

        print("Eventos de los nuevos CREATE:")
        print(f"  ingreso:       {new_events['ingreso']}")
        print(f"  calificación:  {new_events['calificacion']}")
        print(f"  inscripción:   {new_events['inscripcion']}")

        print("BLOCKED restantes:")
        for item in plan["details"]:
            if item["action"] != "BLOCKED":
                continue
            print(
                f"  {item['objective_type']} | {item.get('source_key', '')} | "
                f"{item.get('cop', '')} | filas="
                f"{','.join(map(str, item.get('source_rows') or [item.get('row')]))} | "
                f"{item.get('target_resolution', '')}"
            )

        print(f"Reporte JSON: {json_path}")
        print(f"Plan CSV:     {csv_path}")
        print("No se realizó ninguna escritura en PostgreSQL.")
        return 0

    finally:
        db.rollback()
        db.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (PreflightAbort, base.ImportAbort) as exc:
        print("=== PREFLIGHT 2F-R ABORTADO ===", file=sys.stderr)
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
